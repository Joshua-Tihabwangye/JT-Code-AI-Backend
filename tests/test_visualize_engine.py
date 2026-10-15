"""Every catalog chart and prediction task renders from the sample datasets (engine, in-process)."""

from __future__ import annotations

import base64
import json
from functools import cache

import pytest

from apps.visualize import catalog, samples, viz_engine

LIMITS = {
    "max_bytes": 25 * 1024 * 1024,
    "max_rows": 100_000,
    "max_columns": 200,
    "max_cells": 2_000_000,
    "max_result_bytes": 25 * 1024 * 1024,
    "max_points": 10_000,
    "max_spec_bytes": 5 * 1024 * 1024,
    "preview_rows": 100,
    "allowed_mime_types": ["text/csv"],
}

# chart id -> (sample, fields, options)
CASES: dict[str, tuple[str, dict, dict]] = {
    "kpi_card": ("retail_sales", {"x": "date", "y": "revenue"}, {"time_grain": "M"}),
    "gauge": ("retail_sales", {"y": "revenue", "target": "target_revenue"}, {}),
    "decomposition_tree": ("retail_sales", {"path": ["region", "product", "channel"], "y": "revenue"}, {}),
    "bar": ("retail_sales", {"x": "region", "y": "revenue"}, {"sort": "desc", "show_values": True}),
    "grouped_bar": ("retail_sales", {"x": "region", "y": "revenue", "color": "channel"}, {}),
    "stacked_bar": ("retail_sales", {"x": "product", "y": "units", "color": "region"}, {"orientation": "h"}),
    "percent_stacked_bar": ("retail_sales", {"x": "region", "color": "product"}, {}),
    "lollipop": ("retail_sales", {"x": "product", "y": "profit"}, {}),
    "dot_plot": ("retail_sales", {"x": "region", "y": ["revenue", "cost"]}, {"aggregation": "mean"}),
    "dumbbell": ("retail_sales", {"x": "product", "start": "cost", "end": "revenue"}, {}),
    "radar": (
        "retail_sales",
        {"x": "region", "features": ["units", "revenue", "satisfaction"]},
        {"aggregation": "mean"},
    ),
    "bullet": ("retail_sales", {"x": "region", "y": "revenue", "target": "target_revenue"}, {}),
    "pareto": ("retail_sales", {"x": "product", "y": "revenue"}, {}),
    "category_heatmap": ("retail_sales", {"x": "region", "y": "product", "z": "revenue"}, {}),
    "line": ("retail_sales", {"x": "date", "y": ["revenue", "profit"]}, {"time_grain": "M"}),
    "area": ("website_traffic", {"x": "date", "y": "sessions"}, {}),
    "stacked_area": (
        "website_traffic",
        {"x": "date", "y": "sessions", "color": "source"},
        {"time_grain": "W", "normalize": True},
    ),
    "step_line": ("stock_prices", {"x": "date", "y": "close", "color": "ticker"}, {"aggregation": "mean"}),
    "rolling_average": ("website_traffic", {"x": "date", "y": "conversions"}, {"window": 14}),
    "cumulative": ("retail_sales", {"x": "date", "y": "revenue", "color": "region"}, {"time_grain": "M"}),
    "period_change": ("retail_sales", {"x": "date", "y": "revenue"}, {"time_grain": "M"}),
    "candlestick": (
        "stock_prices",
        {"x": "date", "open": "open", "high": "high", "low": "low", "close": "close"},
        {},
    ),
    "range_band": ("stock_prices", {"x": "date", "y": "close", "lower": "low", "upper": "high"}, {}),
    "calendar_heatmap": ("website_traffic", {"x": "date", "y": "sessions"}, {}),
    "seasonal_decompose": ("retail_sales", {"x": "date", "y": "revenue"}, {"time_grain": "M"}),
    "autocorrelation": ("website_traffic", {"x": "date", "y": "sessions"}, {"time_grain": "D", "lags": 30}),
    "pie": ("retail_sales", {"x": "channel", "y": "revenue"}, {}),
    "donut": ("retail_sales", {"x": "product"}, {}),
    "treemap": ("retail_sales", {"path": ["region", "product"], "y": "revenue"}, {}),
    "sunburst": ("retail_sales", {"path": ["channel", "product"], "y": "units"}, {}),
    "icicle": ("retail_sales", {"path": ["region", "channel", "product"], "y": "revenue"}, {}),
    "waterfall": ("retail_sales", {"x": "product", "y": "profit"}, {}),
    "funnel": ("website_traffic", {"x": "source", "y": "sessions"}, {}),
    "marimekko": ("retail_sales", {"x": "region", "color": "channel", "y": "revenue"}, {}),
    "histogram": ("house_prices", {"x": "price", "color": "neighborhood"}, {"histnorm": "percent"}),
    "density_kde": ("house_prices", {"x": "price", "color": "neighborhood"}, {}),
    "box": ("house_prices", {"y": "price", "x": "neighborhood"}, {"points": "all"}),
    "violin": ("house_prices", {"y": "area_sqft", "x": "bedrooms"}, {}),
    "strip": ("customer_churn", {"y": "monthly_charges", "x": "contract", "color": "churned"}, {}),
    "ecdf": ("customer_churn", {"x": "tenure_months", "color": "churned"}, {}),
    "ridgeline": ("house_prices", {"y": "price", "x": "neighborhood"}, {}),
    "raincloud": ("customer_churn", {"y": "monthly_charges", "x": "contract"}, {}),
    "mean_ci": ("house_prices", {"x": "neighborhood", "y": "price"}, {"confidence": "0.99"}),
    "qq_plot": ("house_prices", {"x": "area_sqft"}, {"distribution": "lognorm"}),
    "distribution_fit": ("house_prices", {"x": "distance_km"}, {"distribution": "gamma"}),
    "control_chart": ("website_traffic", {"x": "date", "y": "bounce_rate"}, {}),
    "bland_altman": ("house_prices", {"x": "price", "y": "predicted_price"}, {}),
    "missing_values": ("retail_sales", {}, {}),
    "scatter": (
        "house_prices",
        {"x": "area_sqft", "y": "price", "color": "neighborhood"},
        {"trendline": "ols", "marginal": "box"},
    ),
    "bubble": (
        "world_indicators",
        {"x": "gdp_per_capita", "y": "life_expectancy", "size": "population_millions", "color": "continent"},
        {},
    ),
    "regression_fit": ("house_prices", {"x": "area_sqft", "y": "price"}, {"degree": 2}),
    "correlation_heatmap": ("house_prices", {}, {"method": "spearman"}),
    "scatter_matrix": (
        "house_prices",
        {"features": ["price", "area_sqft", "age_years"], "color": "neighborhood"},
        {},
    ),
    "density_heatmap": ("house_prices", {"x": "area_sqft", "y": "price"}, {}),
    "density_contour": ("house_prices", {"x": "area_sqft", "y": "price"}, {}),
    "joint_plot": ("house_prices", {"x": "area_sqft", "y": "price"}, {"trendline": "lowess"}),
    "parallel_coordinates": (
        "house_prices",
        {"features": ["price", "area_sqft", "age_years", "distance_km"]},
        {},
    ),
    "parallel_categories": ("customer_churn", {"path": ["contract", "payment_method", "churned"]}, {}),
    "crosstab_heatmap": ("customer_churn", {"x": "churned", "y": "contract"}, {"normalize": "index"}),
    "sankey": ("customer_churn", {"path": ["contract", "payment_method", "churned"]}, {}),
    "network": ("retail_sales", {"source": "region", "target": "product", "y": "revenue"}, {}),
    "choropleth": (
        "world_indicators",
        {"location": "iso3", "y": "life_expectancy"},
        {"location_mode": "ISO-3", "aggregation": "mean"},
    ),
    "scatter_geo": (
        "house_prices",
        {"lat": "lat", "lon": "lon", "color": "neighborhood", "size": "price"},
        {"scope": "africa"},
    ),
    "scatter_3d": (
        "house_prices",
        {"x": "area_sqft", "y": "age_years", "z": "price", "color": "neighborhood"},
        {},
    ),
    "surface_3d": ("house_prices", {"x": "area_sqft", "y": "age_years", "z": "price"}, {}),
    "pca_2d": ("house_prices", {"color": "neighborhood"}, {}),
    "pca_3d": ("house_prices", {}, {}),
    "pca_biplot": ("house_prices", {"color": "neighborhood"}, {}),
    "scree_plot": ("house_prices", {}, {}),
    "tsne_2d": ("world_indicators", {"color": "continent"}, {"perplexity": 8}),
    "lda_2d": (
        "house_prices",
        {"features": ["price", "area_sqft", "distance_km", "age_years"], "color": "neighborhood"},
        {},
    ),
    "kmeans_clusters": (
        "customer_churn",
        {"features": ["tenure_months", "monthly_charges", "support_calls"]},
        {},
    ),
    "confusion_matrix": (
        "customer_churn",
        {"actual": "churned", "predicted": "predicted_churn"},
        {"normalize": "true"},
    ),
    "roc_curve": ("customer_churn", {"actual": "churned", "score": "churn_score"}, {}),
    "pr_curve": ("customer_churn", {"actual": "churned", "score": "churn_score"}, {"positive_class": "Yes"}),
    "calibration_curve": ("customer_churn", {"actual": "churned", "score": "churn_score"}, {}),
    "cumulative_gain": ("customer_churn", {"actual": "churned", "score": "churn_score"}, {}),
    "threshold_metrics": ("customer_churn", {"actual": "churned", "score": "churn_score"}, {}),
    "predicted_vs_actual": ("house_prices", {"actual": "price", "predicted": "predicted_price"}, {}),
    "residual_plot": ("house_prices", {"actual": "price", "predicted": "predicted_price"}, {}),
    "error_distribution": ("house_prices", {"actual": "price", "predicted": "predicted_price"}, {}),
    "feature_importance": (
        "customer_churn",
        {
            "target": "churned",
            "features": ["tenure_months", "monthly_charges", "support_calls", "satisfaction"],
        },
        {},
    ),
    "learning_curve": (
        "house_prices",
        {"target": "price", "features": ["area_sqft", "bedrooms", "age_years", "distance_km"]},
        {"model": "linear"},
    ),
    "elbow": ("customer_churn", {"features": ["tenure_months", "monthly_charges"]}, {"max_k": 6}),
    "silhouette": ("customer_churn", {"features": ["tenure_months", "monthly_charges"]}, {"max_k": 5}),
}


@cache
def _frame(sample_id: str):
    import io

    import pandas as pd

    return pd.read_csv(io.StringIO(samples.sample_csv(sample_id)))


@cache
def _kinds(sample_id: str) -> dict[str, str]:
    return viz_engine.column_kinds(_frame(sample_id))[1]


def test_every_catalog_chart_has_a_renderer_and_a_case():
    ids = {chart.id for chart in catalog.CHARTS}
    assert ids == set(viz_engine.RENDERERS), ids ^ set(viz_engine.RENDERERS)
    assert ids == set(CASES), ids ^ set(CASES)
    assert {chart.category for chart in catalog.CHARTS} == {category.id for category in catalog.CATEGORIES}


HYDRATION = {
    "dataTable",
    "shapes",
    "indexLabel",
    "number",
    "compact",
    "percent",
    "cellValue",
    "labelValue",
    "shapeTooltip",
}


def _assert_hydration_markers_are_known(value):
    """Charts may only reference the fixed client functions (never code)."""
    if isinstance(value, dict):
        if "$fn" in value:
            assert value["$fn"] in HYDRATION, value
        for item in value.values():
            _assert_hydration_markers_are_known(item)
    elif isinstance(value, list):
        for item in value:
            _assert_hydration_markers_are_known(item)


@pytest.mark.parametrize("chart_id", sorted(CASES))
@pytest.mark.parametrize("theme", ["light"])
def test_chart_renders(chart_id: str, theme: str):
    sample, fields, options = CASES[chart_id]
    result = viz_engine.render(
        _frame(sample),
        _kinds(sample),
        {
            "chart_type": chart_id,
            "fields": fields,
            "options": options,
            "style": {"palette": "vivid", "theme": theme},
        },
        LIMITS,
    )
    option = result["option"]
    assert option["series"], f"{chart_id} produced no series"
    assert result["table"]["columns"] and result["table"]["rows"], f"{chart_id} has no organized data table"
    assert result["title"] and result["renderer"] == "echarts"
    explanation = result["explanation"]
    assert explanation["about"] and explanation["howToRead"] and explanation["insight"], explanation
    _assert_hydration_markers_are_known(option)
    json.dumps(result, allow_nan=False, default=str)


def test_dark_theme_and_palette_apply():
    result = viz_engine.render(
        _frame("retail_sales"),
        _kinds("retail_sales"),
        {
            "chart_type": "bar",
            "fields": {"x": "region", "color": "channel"},
            "style": {"palette": "colorblind", "theme": "dark", "title": "Mine"},
        },
        LIMITS,
    )
    option = result["option"]
    assert result["title"] == "Mine"
    assert option["backgroundColor"] == viz_engine.THEME_COLORS["dark"]["bg"]
    assert option["color"][0] == catalog.PALETTES["colorblind"]["colors"][0]


@pytest.mark.parametrize(
    ("chart_id", "fields", "message"),
    [
        ("bar", {}, "Choose a column"),
        ("bar", {"x": "nope"}, "Unknown column"),
        ("histogram", {"x": "region"}, "needs a numeric column"),
        ("bar", {"x": "region", "evil": "x"}, "Unknown fields"),
    ],
)
def test_invalid_selections_are_rejected(chart_id, fields, message):
    with pytest.raises(viz_engine.AnalysisError, match=message):
        viz_engine.render(
            _frame("retail_sales"), _kinds("retail_sales"), {"chart_type": chart_id, "fields": fields}, LIMITS
        )


PREDICTIONS = [
    ("trend", "retail_sales", {"date": "date", "target": "revenue"}, "robust", {"time_grain": "M"}),
    (
        "trend",
        "website_traffic",
        {"date": "date", "target": "sessions"},
        "robust",
        {"time_grain": "D", "window": 7},
    ),
    (
        "drivers",
        "customer_churn",
        {"target": "churned", "features": ["contract", "tenure_months", "support_calls", "payment_method"]},
        "segments",
        {},
    ),
    ("drivers", "house_prices", {"target": "price"}, "segments", {}),
    (
        "forecast",
        "retail_sales",
        {"date": "date", "target": "revenue", "split_by": "region"},
        "holt_winters",
        {"horizon": 6, "time_grain": "M"},
    ),
    (
        "forecast",
        "retail_sales",
        {"date": "date", "target": "revenue"},
        "auto",
        {"horizon": 6, "time_grain": "M"},
    ),
    (
        "forecast",
        "website_traffic",
        {"date": "date", "target": "sessions"},
        "holt_winters",
        {"horizon": 14, "time_grain": "D"},
    ),
    (
        "forecast",
        "retail_sales",
        {"date": "date", "target": "units"},
        "arima",
        {"horizon": 4, "time_grain": "W"},
    ),
    (
        "regression",
        "house_prices",
        {"target": "price", "features": ["area_sqft", "bedrooms", "neighborhood", "age_years"]},
        "auto",
        {},
    ),
    ("regression", "house_prices", {"target": "price"}, "ridge", {}),
    (
        "classification",
        "customer_churn",
        {"target": "churned", "features": ["tenure_months", "contract", "monthly_charges", "support_calls"]},
        "auto",
        {},
    ),
    (
        "classification",
        "world_indicators",
        {"target": "continent", "features": ["gdp_per_capita", "life_expectancy", "co2_per_capita"]},
        "random_forest",
        {"test_size": "0.3"},
    ),
    (
        "clustering",
        "customer_churn",
        {"features": ["tenure_months", "monthly_charges", "support_calls"]},
        "kmeans",
        {},
    ),
    ("clustering", "house_prices", {"features": ["price", "area_sqft"]}, "gaussian_mixture", {"clusters": 3}),
    ("anomaly", "house_prices", {"features": ["price", "area_sqft", "distance_km"]}, "isolation_forest", {}),
    ("anomaly", "retail_sales", {"features": ["revenue", "units"]}, "zscore", {"contamination": "0.02"}),
    ("anomaly", "house_prices", {"features": ["price", "area_sqft"]}, "lof", {}),
]


@pytest.mark.parametrize(("task", "sample", "fields", "model", "options"), PREDICTIONS)
def test_prediction_tasks(task, sample, fields, model, options):
    result = viz_engine.predict(
        _frame(sample),
        _kinds(sample),
        {"task": task, "fields": fields, "model": model, "options": options},
        LIMITS,
    )
    assert result["metrics"] and result["charts"] and result["summary"]
    assert all(chart["option"]["series"] and chart["explanation"]["howToRead"] for chart in result["charts"])
    assert all(chart["table"]["columns"] for chart in result["charts"])
    assert base64.b64decode(result["result"])
    json.dumps(viz_engine._clean_nan(result), allow_nan=False, default=str)


def test_insights_are_deterministic_and_suggest_charts():
    frame = _frame("retail_sales")
    profile, kinds = viz_engine.column_kinds(frame)
    first = viz_engine.insights(viz_engine.typed(frame, kinds), profile, kinds)
    second = viz_engine.insights(viz_engine.typed(frame, kinds), profile, kinds)
    assert first == second
    assert any(item["kind"] == "correlation" for item in first)
    for item in first:
        if "chart" in item:
            assert item["chart"]["chartType"] in catalog.CHARTS_BY_ID


def test_import_formats_and_organize():
    import io

    import pandas as pd

    frame = pd.DataFrame(
        {"Region ": ["East", "West", None, "East"], "Sales": ["1,200", "800", "500", "1,200"]}
    )
    excel = io.BytesIO()
    frame.to_excel(excel, index=False)
    sources = {
        "xlsx": excel.getvalue(),
        "json": frame.to_json(orient="records").encode(),
        "tsv": frame.to_csv(sep="\t", index=False).encode(),
        "txt": frame.to_csv(sep=";", index=False).encode(),
    }
    for fmt, content in sources.items():
        imported = viz_engine.import_table(
            {"format": fmt, "data": base64.b64encode(content).decode()}, LIMITS
        )
        assert list(imported.columns) == ["Region", "Sales"], fmt
    imported = viz_engine.import_table(
        {"format": "json", "data": base64.b64encode(sources["json"]).decode()}, LIMITS
    )
    kinds = viz_engine.column_kinds(imported)[1]
    organized = viz_engine.organize(
        imported,
        kinds,
        {
            "clean": {
                "cast": {"Sales": "number"},
                "fill_missing": {"Region": {"strategy": "value", "value": "Unknown"}},
                "dedupe": True,
                "rename": {"Sales": "Revenue"},
            },
            "transform": {"sort": {"column": "Sales", "direction": "desc"}},
        },
    )
    assert list(organized.columns) == ["Region", "Revenue"]
    assert organized["Revenue"].tolist() == [1200, 800, 500]
    assert "Unknown" in organized["Region"].tolist()
    with pytest.raises(viz_engine.AnalysisError):
        viz_engine.organize(imported, kinds, {"clean": {"eval": "os.system('x')"}})
    with pytest.raises(viz_engine.AnalysisError):
        viz_engine.import_table({"format": "xlsx", "data": base64.b64encode(b"not a zip").decode()}, LIMITS)


def test_xlsx_export_neutralizes_formulas():
    import io

    import pandas as pd
    from openpyxl import load_workbook

    frame = pd.DataFrame({"name": ['=HYPERLINK("http://x")', "ok"], "value": [1, 2]})
    profile, kinds = viz_engine.column_kinds(frame)
    exported = viz_engine.export_dataset(frame, kinds, profile, {"format": "xlsx", "name": "t"}, LIMITS)
    book = load_workbook(io.BytesIO(base64.b64decode(exported["content"])))
    assert book.sheetnames[:3] == ["Summary", "Data", "Profile"]
    assert book["Data"]["A2"].value.startswith("'=")


def test_custom_colours_override_the_palette():
    result = viz_engine.render(
        _frame("retail_sales"),
        _kinds("retail_sales"),
        {
            "chart_type": "bar",
            "fields": {"x": "region", "color": "channel"},
            "style": {"colors": ["#112233", "#445566", "bad"]},
        },
        LIMITS,
    )
    assert result["option"]["color"][:2] == ["#112233", "#445566"]
    assert "bad" not in result["option"]["color"]


def test_organized_tables_follow_the_chart():
    bar = viz_engine.render(
        _frame("retail_sales"),
        _kinds("retail_sales"),
        {"chart_type": "grouped_bar", "fields": {"x": "region", "y": "revenue", "color": "channel"}},
        LIMITS,
    )
    assert bar["table"]["columns"][0] == "Region" and set(bar["table"]["columns"][1:]) == {
        "Online",
        "Retail store",
        "Partner",
    }
    pie = viz_engine.render(
        _frame("retail_sales"),
        _kinds("retail_sales"),
        {"chart_type": "pie", "fields": {"x": "channel"}},
        LIMITS,
    )
    assert pie["table"]["columns"] == ["Name", "Count", "Share"]
    assert sum(float(row[2].rstrip("%")) for row in pie["table"]["rows"]) == pytest.approx(100, abs=0.5)


def test_column_names_read_like_words():
    assert viz_engine.c.nice("unit_price") == "Unit price"
    assert viz_engine.c.nice("orderId") == "Order ID"
    assert viz_engine.c.nice("customer_id") == "Customer ID"
    line = viz_engine.render(
        _frame("retail_sales"),
        _kinds("retail_sales"),
        {
            "chart_type": "line",
            "fields": {"x": "date", "y": "target_revenue"},
            "options": {"time_grain": "M"},
        },
        LIMITS,
    )
    text = json.dumps([line["title"], line["explanation"], line["stats"], line["table"]["columns"]])
    assert "target_revenue" not in text and "target revenue" in line["title"]
    heat = viz_engine.render(
        _frame("house_prices"),
        _kinds("house_prices"),
        {"chart_type": "correlation_heatmap", "fields": {}, "options": {}},
        LIMITS,
    )
    assert "Area sqft" in json.dumps(heat["option"]) and "area_sqft" not in json.dumps(heat["table"])
    forecast = viz_engine.predict(
        _frame("retail_sales"),
        _kinds("retail_sales"),
        {
            "task": "forecast",
            "fields": {"date": "date", "target": "target_revenue"},
            "model": "holt_winters",
            "options": {"horizon": 3},
        },
        LIMITS,
    )
    shown = json.dumps([forecast["summary"], forecast["metrics"], [c["title"] for c in forecast["charts"]]])
    assert "target_revenue" not in shown


def _workbook(charts, graphs=True, **extra):
    import io

    import openpyxl

    result = viz_engine.report_workbook(
        {"charts": charts, "summary": [("Report", "Test")], "graphs": graphs, **extra}, LIMITS
    )
    return openpyxl.load_workbook(io.BytesIO(base64.b64decode(result["content"])))


def test_excel_workbooks_carry_values_and_native_charts():
    rendered = []
    for chart_type, sample, fields in (
        ("bar", "retail_sales", {"x": "region", "y": "revenue"}),
        ("donut", "retail_sales", {"x": "product"}),
        ("sankey", "customer_churn", {"path": ["contract", "payment_method", "churned"]}),
    ):
        out = viz_engine.render(
            _frame(sample), _kinds(sample), {"chart_type": chart_type, "fields": fields}, LIMITS
        )
        rendered.append(
            {
                "title": out["title"],
                "subtitle": chart_type,
                "option": out["option"],
                "stats": out["stats"][:1],
            }
        )
    book = _workbook(rendered)
    assert book.sheetnames[0] == "Summary" and book.sheetnames[1].startswith("Chart 1")
    bar_sheet, donut_sheet, sankey_sheet = (book[name] for name in book.sheetnames[1:4])
    assert len(bar_sheet._charts) == 1 and len(donut_sheet._charts) == 1
    assert not sankey_sheet._charts and sankey_sheet.max_row > 5  # values only: no Excel equivalent
    assert bar_sheet["A5"].value == "Region" and isinstance(bar_sheet["B6"].value, float)
    links = [cell.hyperlink.location for row in book["Summary"].iter_rows() for cell in row if cell.hyperlink]
    assert len(links) == 3 and links[0].startswith("'Chart 1")
    values_only = _workbook(rendered, graphs=False)
    assert not any(values_only[name]._charts for name in values_only.sheetnames)
    assert values_only[values_only.sheetnames[1]].max_row > 5


def test_excel_workbook_renders_chart_specs_and_tables():
    data = base64.b64encode(samples.sample_csv("retail_sales").encode()).decode()
    book = _workbook(
        [
            {
                "dataset": 0,
                "chart_type": "line",
                "fields": {"x": "date", "y": "revenue"},
                "options": {"time_grain": "M"},
            },
            {"dataset": 0, "chart_type": "bar", "fields": {"x": "no_such_column"}},
        ],
        datasets=[{"name": "Retail", "data": data}],
        tables=[{"name": "Predictions", "data": base64.b64encode(b"a,b\n1,2\n").decode()}],
    )
    assert book.sheetnames[1].startswith("Chart 1") and "Predictions" in book.sheetnames
    assert "Retail data" in book.sheetnames and "Retail profile" in book.sheetnames
    assert len(book[book.sheetnames[1]]._charts) == 1
    summary = " ".join(str(cell.value) for row in book["Summary"].iter_rows() for cell in row)
    assert "not included" in summary


def test_a_bar_with_only_a_category_counts_rows():
    """Studio: Category = Channel and Value = None must draw counts (it once rendered blank)."""
    out = viz_engine.render(
        _frame("retail_sales"),
        _kinds("retail_sales"),
        {"chart_type": "bar", "fields": {"x": "channel"}},
        LIMITS,
    )
    assert out["title"] == "Count by channel"
    assert sorted(row[1] for row in out["table"]["rows"]) == [375, 869, 1156]
