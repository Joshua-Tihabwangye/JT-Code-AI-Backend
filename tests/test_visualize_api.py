"""Visualization API: end-to-end flow through the real sandbox, metering, assistant and access control."""

from __future__ import annotations

import io
import json
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient

from apps.billing.models import CreditWallet
from apps.identity.models import Organization, Role, User, UserOrganization, UserRole


def _org(label: str) -> tuple[Organization, User]:
    owner = User.objects.create_user(
        username=f"{label}-owner", supabase_user_id=f"sb-{label}", email=f"{label}@example.test"
    )
    org = Organization.objects.create(name=f"{label} org", slug=f"{label}-org", owner=owner)
    UserOrganization.objects.get_or_create(user=owner, organization=org)
    CreditWallet.objects.update_or_create(organization=org, defaults={"balance": Decimal("1000")})
    return org, owner


def _member(org: Organization, label: str, role: str) -> User:
    user = User.objects.create_user(
        username=f"{label}-{role}",
        supabase_user_id=f"sb-{label}-{role}",
        email=f"{label}-{role}@example.test",
    )
    UserOrganization.objects.get_or_create(user=user, organization=org)
    UserRole.objects.filter(user=user, organization=org).delete()
    UserRole.objects.create(user=user, organization=org, role=Role.objects.get_or_create(name=role)[0])
    return user


def _client(user: User, org: Organization) -> APIClient:
    api = APIClient()
    api.force_authenticate(user)
    api.credentials(HTTP_X_ORGANIZATION_ID=str(org.id))
    return api


def _sample(api: APIClient, sample: str = "retail_sales") -> dict:
    response = api.post("/api/v1/visualize/samples/", {"sampleId": sample}, format="json")
    assert response.status_code == 201, response.content
    return response.json()


@pytest.mark.django_db
def test_catalog_lists_every_category_chart_task_and_sample():
    org, owner = _org("catalog")
    body = _client(owner, org).get("/api/v1/visualize/catalog/").json()
    assert len(body["categories"]) == 10
    assert len(body["charts"]) >= 80
    assert {task["id"] for task in body["predictionTasks"]} == {
        "forecast",
        "regression",
        "classification",
        "clustering",
        "anomaly",
        "trend",
        "drivers",
    }
    assert {sample["id"] for sample in body["samples"]} >= {"retail_sales", "customer_churn"}
    assert body["pricing"]["render"] == 0
    assert APIClient().get("/api/v1/visualize/catalog/").status_code in {401, 403}


@pytest.mark.django_db
def test_dataset_to_chart_flow():
    org, owner = _org("flow")
    api = _client(owner, org)
    dataset = _sample(api)

    preview = api.get(f"/api/v1/visualize/datasets/{dataset['id']}/preview/?rows=5").json()
    assert len(preview["rows"]) == 5 and preview["total"] == 2400

    insights = api.get(f"/api/v1/visualize/datasets/{dataset['id']}/insights/").json()
    assert insights["insights"] and insights["recommendations"]
    recommended = api.get(f"/api/v1/visualize/datasets/{dataset['id']}/recommend/?chart=scatter").json()
    assert set(recommended["fields"]) >= {"x", "y"}

    render = api.post(
        "/api/v1/visualize/render/",
        {
            "datasetId": dataset["id"],
            "chartType": "bar",
            "fields": {"x": "region", "y": "revenue"},
            "options": {"sort": "desc"},
            "style": {"palette": "sunset", "theme": "dark"},
        },
        format="json",
    )
    assert render.status_code == 200, render.content
    rendered = render.json()
    assert rendered["renderer"] == "echarts" and rendered["option"]["series"]
    assert rendered["explanation"]["howToRead"] and rendered["explanation"]["insight"]
    bad = api.post(
        "/api/v1/visualize/render/",
        {"datasetId": dataset["id"], "chartType": "histogram", "fields": {"x": "region"}},
        format="json",
    )
    assert bad.status_code == 400 and "numeric" in bad.json()["message"]

    chart = api.post(
        "/api/v1/visualize/charts/",
        {
            "datasetId": dataset["id"],
            "chartType": "line",
            "fields": {"x": "date", "y": ["revenue"]},
            "title": "<script>alert(1)</script> Revenue",
        },
        format="json",
    )
    assert chart.status_code == 201, chart.content
    chart_id = chart.json()["id"]
    assert api.get("/api/v1/visualize/charts/").json()[
        "results" if "results" in api.get("/api/v1/visualize/charts/").json() else 0
    ]

    page = api.get(f"/api/v1/visualize/charts/{chart_id}/export/?type=html")
    assert page.status_code == 200
    text = page.content.decode()
    assert "echarts.init" in text and "jtHydrate" in text and "Apache Software Foundation" in text
    assert "How to read it" in text and "Key takeaway" in text
    assert "<script>alert(1)</script>" not in text and "&lt;script&gt;alert(1)&lt;/script&gt;" in text
    from openpyxl import load_workbook

    exported = api.get(f"/api/v1/visualize/datasets/{dataset['id']}/export/?type=xlsx")
    assert exported.status_code == 200 and exported["Content-Disposition"].endswith('.xlsx"')
    assert load_workbook(io.BytesIO(exported.content)).sheetnames == ["Summary", "Data", "Profile"]
    with_graphs = load_workbook(
        io.BytesIO(api.get(f"/api/v1/visualize/datasets/{dataset['id']}/export/?type=xlsx&graphs=1").content)
    )
    chart_sheets = [name for name in with_graphs.sheetnames if name.startswith("Chart")]
    assert len(chart_sheets) >= 3 and sum(len(with_graphs[name]._charts) for name in chart_sheets) >= 3

    studio = api.post(
        "/api/v1/visualize/workbook/",
        {
            "title": "My page",
            "charts": [
                {"datasetId": dataset["id"], "chartType": "bar", "fields": {"x": "region", "y": "revenue"}},
                {
                    "datasetId": dataset["id"],
                    "chartType": "donut",
                    "fields": {"x": "channel"},
                    "style": {"colors": ["#112233"]},
                    "title": "Channel mix",
                },
            ],
        },
        format="json",
    )
    assert studio.status_code == 200, studio.content
    studio_book = load_workbook(io.BytesIO(studio.content))
    assert studio_book.sheetnames[2].startswith("Chart 2 Channel mix")
    assert all(len(studio_book[name]._charts) == 1 for name in studio_book.sheetnames[1:3])
    too_many = api.post(
        "/api/v1/visualize/workbook/",
        {"charts": [{"datasetId": dataset["id"], "chartType": "bar"}] * 13},
        format="json",
    )
    assert too_many.status_code == 400
    other_org, other = _org("flow-other")
    hidden = _client(other, other_org).post(
        "/api/v1/visualize/workbook/",
        {"charts": [{"datasetId": dataset["id"], "chartType": "bar", "fields": {"x": "region"}}]},
        format="json",
    )
    assert hidden.status_code == 404
    assert api.get(f"/api/v1/visualize/charts/{chart_id}/export/?type=html").status_code == 200


@pytest.mark.django_db
def test_import_and_organize():
    org, owner = _org("import")
    api = _client(owner, org)
    import pandas as pd

    frame = pd.DataFrame({"Region": ["East", "West", "East", None], "Sales": [100, 250, 100, 40]})
    excel = io.BytesIO()
    frame.to_excel(excel, index=False)
    upload = SimpleUploadedFile("sales.xlsx", excel.getvalue(), content_type="application/octet-stream")
    created = api.post("/api/v1/visualize/datasets/import/", {"file": upload}, format="multipart")
    assert created.status_code == 201, created.content
    assert created.json()["name"] == "sales" and created.json()["row_count"] == 4
    unsupported = SimpleUploadedFile("evil.exe", b"MZ", content_type="application/octet-stream")
    assert (
        api.post("/api/v1/visualize/datasets/import/", {"file": unsupported}, format="multipart").status_code
        == 400
    )

    steps = {
        "clean": {"dedupe": True, "fill_missing": {"Region": "mode"}},
        "transform": {"group_by": "Region", "aggregate": {"Sales": "sum"}},
    }
    preview = api.post(
        f"/api/v1/visualize/datasets/{created.json()['id']}/organize/", {"steps": steps}, format="json"
    )
    assert preview.status_code == 200, preview.content
    assert {row["Region"]: row["Sales_sum"] for row in preview.json()["preview"]} == {
        "East": 140,
        "West": 250,
    }
    saved = api.post(
        f"/api/v1/visualize/datasets/{created.json()['id']}/organize/",
        {"steps": steps, "saveAs": "By region"},
        format="json",
    )
    assert saved.status_code == 201 and saved.json()["name"] == "By region"
    rejected = api.post(
        f"/api/v1/visualize/datasets/{created.json()['id']}/organize/",
        {"steps": {"python": "import os"}},
        format="json",
    )
    assert rejected.status_code == 400


@pytest.mark.django_db
def test_prediction_is_metered_and_results_can_be_saved():
    org, owner = _org("predict")
    api = _client(owner, org)
    dataset = _sample(api, "customer_churn")
    before = CreditWallet.objects.get(organization=org).balance
    response = api.post(
        "/api/v1/visualize/predict/",
        {
            "datasetId": dataset["id"],
            "task": "classification",
            "model": "logistic",
            "fields": {"target": "churned", "features": ["tenure_months", "contract", "monthly_charges"]},
        },
        format="json",
    )
    assert response.status_code == 200, response.content
    body = response.json()
    assert any(metric["key"] == "accuracy" for metric in body["metrics"]) and body["charts"]
    assert CreditWallet.objects.get(organization=org).balance == before - Decimal("10")

    download = api.get(f"/api/v1/visualize/predict/results/{body['resultId']}/")
    assert download.status_code == 200 and b"predicted_churned" in download.content
    from openpyxl import load_workbook

    excel = api.get(f"/api/v1/visualize/predict/results/{body['resultId']}/?type=xlsx")
    assert excel.status_code == 200, excel.content
    book = load_workbook(io.BytesIO(excel.content))
    assert "Predictions" in book.sheetnames and book.sheetnames[1].startswith("Chart 1")
    assert any(book[name]._charts for name in book.sheetnames if name.startswith("Chart"))
    summary = {row[0]: row[1] for row in book["Summary"].iter_rows(values_only=True)}
    assert summary["Prediction"] and summary["Model"], summary
    assert any("ccuracy" in str(key) for key in summary), summary
    saved = api.post(
        f"/api/v1/visualize/predict/results/{body['resultId']}/", {"name": "Churn scored"}, format="json"
    )
    assert saved.status_code == 201
    chart = api.post(
        "/api/v1/visualize/charts/", {"resultId": body["resultId"], "chartIndex": 0}, format="json"
    )
    assert chart.status_code == 201 and chart.json()["chartType"].startswith("prediction:")

    failed = api.post(
        "/api/v1/visualize/predict/",
        {
            "datasetId": dataset["id"],
            "task": "forecast",
            "fields": {"date": "churned", "target": "tenure_months"},
        },
        format="json",
    )
    assert failed.status_code == 400
    assert CreditWallet.objects.get(organization=org).balance == before - Decimal("10"), (
        "failed runs are not charged"
    )
    other_org, other = _org("predict-other")
    assert (
        _client(other, other_org).get(f"/api/v1/visualize/predict/results/{body['resultId']}/").status_code
        == 404
    )


@pytest.mark.django_db
def test_assistant_uses_tokens_only_on_prompt_and_validates_the_model(monkeypatch):
    from apps.ai_gateway import service

    org, owner = _org("assist")
    api = _client(owner, org)
    dataset = _sample(api)
    calls = []

    class Outcome:
        content = json.dumps(
            {
                "reply": "A stacked bar shows the channel mix.",
                "action": {
                    "kind": "chart",
                    "chartType": "stacked_bar",
                    "fields": {"x": "region", "color": "channel", "y": "made_up"},
                    "options": {"aggregation": "drop table"},
                },
            }
        )
        usage = type("U", (), {"input_tokens": 420, "output_tokens": 60})()
        run = type("R", (), {"provider_cost_usd": Decimal("0.0004")})()
        model = type("M", (), {"name": "test-model"})()

    def fake(**kwargs):
        calls.append(kwargs)
        return Outcome()

    monkeypatch.setattr(service, "generate_completion", fake)
    # Rendering, insights and recommendations never touch the model.
    api.get(f"/api/v1/visualize/datasets/{dataset['id']}/insights/")
    api.post(
        "/api/v1/visualize/render/",
        {"datasetId": dataset["id"], "chartType": "pie", "fields": {"x": "channel"}},
        format="json",
    )
    assert calls == []

    answer = api.post(
        "/api/v1/visualize/assistant/",
        {"prompt": "Show the channel mix per region", "datasetId": dataset["id"]},
        format="json",
    ).json()
    assert len(calls) == 1
    sent = calls[0]["messages"][1].content
    assert "Show the channel mix per region" in sent and "East" in sent  # column summary, not raw rows
    assert "2024-0" not in sent.split("<dataset>")[1].split("</dataset>")[0].replace("2024-01-01..", "")
    assert answer["source"] == "ai"
    assert answer["action"]["chartType"] == "stacked_bar"
    assert answer["action"]["fields"]["x"] == "region" and answer["action"]["fields"].get("y") != "made_up"
    assert "aggregation" not in answer["action"]["options"]
    assert Decimal(answer["usage"]["credits"]) > 0

    from apps.ai_gateway.adapters import AIGatewayError

    def broken(**kwargs):
        raise AIGatewayError("down")

    monkeypatch.setattr(service, "generate_completion", broken)
    before = CreditWallet.objects.get(organization=org).balance
    fallback = api.post(
        "/api/v1/visualize/assistant/",
        {"prompt": "forecast revenue for next month", "datasetId": dataset["id"]},
        format="json",
    ).json()
    assert fallback["source"] == "heuristic" and fallback["action"]["kind"] == "predict"
    assert fallback["action"]["task"] == "forecast"
    assert CreditWallet.objects.get(organization=org).balance == before


@pytest.mark.django_db
def test_access_control():
    org, owner = _org("acl")
    api = _client(owner, org)
    dataset = _sample(api)
    chart = api.post(
        "/api/v1/visualize/charts/",
        {"datasetId": dataset["id"], "chartType": "pie", "fields": {"x": "channel"}},
        format="json",
    ).json()

    other_org, other = _org("acl-other")
    stranger = _client(other, other_org)
    assert stranger.get(f"/api/v1/visualize/datasets/{dataset['id']}/preview/").status_code == 404
    assert (
        stranger.post(
            "/api/v1/visualize/render/",
            {"datasetId": dataset["id"], "chartType": "pie", "fields": {"x": "channel"}},
            format="json",
        ).status_code
        == 404
    )
    assert stranger.get(f"/api/v1/visualize/charts/{chart['id']}/").status_code == 404

    viewer = _client(_member(org, "acl", Role.RoleType.VIEWER), org)
    assert (
        viewer.post("/api/v1/visualize/samples/", {"sampleId": "retail_sales"}, format="json").status_code
        == 403
    )
    assert viewer.get(f"/api/v1/visualize/charts/{chart['id']}/").status_code == 404, (
        "charts are private until shared"
    )
    api.patch(f"/api/v1/visualize/charts/{chart['id']}/", {"isShared": True}, format="json")
    assert viewer.get(f"/api/v1/visualize/charts/{chart['id']}/").status_code == 200
    assert viewer.delete(f"/api/v1/visualize/charts/{chart['id']}/").status_code == 403


@pytest.mark.django_db
def test_quality_scan_flags_problems_without_changing_the_data():
    org, owner = _org("quality")
    api = _client(owner, org)
    csv = (
        "order_id,customer,email,amount,region,order_date\n"
        "1,Ann,ann@example.com,120,East,2025-01-02\n"
        "2,Bob,bob@example,95,east ,2025-01-03\n"
        "2,Cy,cy@example.com,-40,East,2025-01-04\n"
        ",Dee,dee@example.com,N/A,West,2025-01-05\n"
        "5,Eve,eve@example.com,110,West,2199-01-01\n"
        "6,Fay,fay@example.com,105,West,2025-01-07\n"
        "6,Fay,fay@example.com,105,West,2025-01-07\n"
        "8,Gus,gus@example.com,99000,North,2025-01-08\n"
        "9,Hal,hal@example.com,101,North,2025-01-09\n"
        "10,Ida,ida@example.com,97,North,2025-01-10\n"
        "11,Jo,jo@example.com,103,North,2025-01-11\n"
        "12,Kay,kay@example.com,99,North,2025-01-12\n"
    )
    dataset = api.post(
        "/api/v1/analysis/datasets/", {"name": "Orders", "inline_data": csv}, format="json"
    ).json()
    report = api.get(f"/api/v1/visualize/datasets/{dataset['id']}/quality/").json()
    found = {issue["type"] for issue in report["issues"]}
    assert {
        "duplicate_row",
        "duplicate_id",
        "missing_id",
        "id_gap",
        "placeholder",
        "invalid_email",
        "negative",
        "inconsistent_category",
        "whitespace",
        "invalid_date",
        "outlier",
    } <= found, found
    assert 0 < report["summary"]["score"] < 100
    gap = next(issue for issue in report["issues"] if issue["type"] == "id_gap")
    assert {"3", "4", "7"} <= set(gap["examples"])

    page = api.get(f"/api/v1/visualize/datasets/{dataset['id']}/preview/?issues=1&filter=issues").json()
    assert page["total"] < page["datasetRows"] == 12
    assert "duplicate_row" in page["flags"]["6"]["*"]
    assert "placeholder" in page["flags"]["3"]["amount"]
    unchanged = api.get(f"/api/v1/visualize/datasets/{dataset['id']}/preview/?rows=50").json()
    assert unchanged["total"] == 12, "scanning never edits the data"

    # The person decides what to fix: edit cells and delete the duplicate row.
    saved = api.post(
        f"/api/v1/visualize/datasets/{dataset['id']}/edit/",
        {
            "edits": [
                {"row": 3, "column": "order_id", "value": "4"},
                {"row": 3, "column": "amount", "value": "88"},
            ],
            "deleteRows": [6],
        },
        format="json",
    )
    assert saved.status_code == 200, saved.content
    after = api.get(f"/api/v1/visualize/datasets/{dataset['id']}/quality/").json()
    assert not {"duplicate_row", "missing_id", "placeholder"} & {issue["type"] for issue in after["issues"]}

    viewer = _client(_member(org, "quality", Role.RoleType.EDITOR), org)
    blocked = viewer.post(
        f"/api/v1/visualize/datasets/{dataset['id']}/edit/", {"deleteRows": [0]}, format="json"
    )
    assert blocked.status_code in {403, 404}

    fixed = SimpleUploadedFile("orders.csv", b"order_id,amount\n1,10\n2,20\n", content_type="text/csv")
    replaced = api.post(
        f"/api/v1/visualize/datasets/{dataset['id']}/replace/", {"file": fixed}, format="multipart"
    )
    assert replaced.status_code == 200 and replaced.json()["row_count"] == 2


@pytest.mark.django_db
def test_world_map_is_served_for_map_charts():
    org, owner = _org("maps")
    response = _client(owner, org).get("/api/v1/visualize/maps/world/")
    assert response.status_code == 200
    world = json.loads(response.content)
    assert world["type"] == "FeatureCollection" and len(world["features"]) > 150
