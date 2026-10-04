"""Content extraction for RAG ingestion.

The first stage of the Agentic RAG pipeline turns a knowledge source into
normalized UTF-8 text that can be chunked and embedded:

* ``text`` sources embed their configured content directly;
* ``url`` sources are fetched through the egress policy (HTML, PDF, text/JSON);
* ``file`` sources read a registered ImageKit asset through a short-lived signed
  URL (PDF, DOCX, HTML, Markdown, CSV, JSON and plain text).

PDF extraction records each page's start offset so chunks carry page numbers.
Output is canonicalized to ``\\n`` line endings and bounded in size.
"""

from __future__ import annotations

import html
import json
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO
from typing import Any

from django.conf import settings

from apps.tools.egress import EgressDenied, EgressError, safe_request

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_TEXT_TYPES = {"application/json", "application/xml", "application/xhtml+xml", "application/csv"}
_TEXT_EXTENSIONS = {"txt", "md", "markdown", "csv", "json", "xml"}


class ExtractionError(RuntimeError):
    """Raised when a knowledge source cannot be converted to text."""


@dataclass(frozen=True)
class ExtractedText:
    text: str
    mime_type: str
    page_offsets: list[int] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.page_offsets)


class _TextHTMLParser(HTMLParser):
    """Visible-text extraction that keeps block structure and drops scripts."""

    _SKIP = {"script", "style", "noscript", "template", "svg", "head"}
    _BLOCK = {
        "p", "div", "section", "article", "header", "footer", "li", "ul", "ol", "table", "tr",
        "br", "hr", "pre", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6",
    }  # fmt: skip

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in {"h1", "h2", "h3", "h4", "h5", "h6"} and not self._skip_depth:
            # Preserve headings as Markdown so chunks keep their section path.
            self._parts.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag == "li" and not self._skip_depth:
            self._parts.append("\n- ")
        elif tag in self._BLOCK:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK and tag != "li":
            self._parts.append("\n\n" if tag in {"p", "h1", "h2", "h3", "h4", "h5", "h6"} else "\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth:
            self._parts.append(data)

    def text(self) -> str:
        lines = (" ".join(line.split()) for line in "".join(self._parts).split("\n"))
        output: list[str] = []
        for line in lines:
            if line or (output and output[-1]):
                output.append(line)
        return "\n".join(output).strip()


def _clean_html(markup: str) -> str:
    parser = _TextHTMLParser()
    parser.feed(markup)
    parser.close()
    return html.unescape(parser.text())


def _normalize(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _cap_size(text: str) -> str:
    encoded = len(text.encode("utf-8"))
    limit = settings.RAG_MAX_EXTRACTED_BYTES
    if encoded > limit:
        raise ExtractionError(
            f"Source content is {encoded} bytes, exceeding the {limit}-byte RAG extraction limit."
        )
    return text


def _result(text: str, mime_type: str, page_offsets: list[int] | None = None) -> ExtractedText:
    return ExtractedText(
        text=_cap_size(_normalize(text)), mime_type=mime_type, page_offsets=page_offsets or []
    )


def extract_source_text(
    *, source_type: str, config: dict[str, Any] | None = None, metadata: dict[str, Any] | None = None
) -> ExtractedText:
    """Extract normalized text from a knowledge source.

    ``config`` mirrors ``Source.config``; ``metadata`` mirrors ``Document.metadata``.
    """
    config = config or {}
    metadata = metadata or {}
    source_type = (source_type or "").lower()

    if source_type == "text":
        text = config.get("text") or ""
        if not isinstance(text, str) or not text.strip():
            raise ExtractionError("TEXT source has no content in its configuration.")
        return _result(text, str(config.get("mime_type") or "text/plain"))

    if source_type == "url":
        return _fetch_url(str(config.get("url") or ""))

    if source_type == "integration":
        # Pushed by the knowledge-integration-sync n8n workflow (apps.orchestration.knowledge).
        pushed = metadata.get("integration") or {}
        if pushed.get("text"):
            return _result(str(pushed["text"]), "text/plain")
        if pushed.get("url"):
            return _fetch_url(str(pushed["url"]))
        raise ExtractionError("The integration document has no content.")

    if source_type == "file":
        asset_id = config.get("asset_id")
        if not asset_id:
            raise ExtractionError("FILE sources require config.asset_id.")
        return _extract_asset_text(str(asset_id), metadata)

    raise ExtractionError(f"Extraction for source_type={source_type!r} is not supported.")


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("Text content must be UTF-8 encoded.") from exc


def _from_bytes(content: bytes, *, content_type: str, extension: str = "") -> ExtractedText:
    content_type = content_type.split(";", 1)[0].strip().lower()
    extension = extension.lower().lstrip(".")
    if "pdf" in content_type or extension == "pdf":
        return _extract_pdf(content)
    if content_type == _DOCX or extension == "docx":
        return _extract_docx(content)
    if "html" in content_type or extension in {"html", "htm"}:
        return _result(_clean_html(_decode(content)), "text/html")
    if content_type == "application/json" or extension == "json":
        try:
            parsed = json.loads(_decode(content))
        except ValueError as exc:
            raise ExtractionError("JSON content could not be parsed.") from exc
        return _result(json.dumps(parsed, indent=2, ensure_ascii=False), "application/json")
    if content_type.startswith("text/") or content_type in _TEXT_TYPES or extension in _TEXT_EXTENSIONS:
        return _result(_decode(content), content_type or "text/plain")
    raise ExtractionError(f"Unsupported content type: {content_type or extension or 'unknown'}.")


def _fetch_url(url: str) -> ExtractedText:
    if not url:
        raise ExtractionError("URL source has no address in its configuration.")
    try:
        response = safe_request(
            "GET",
            url,
            headers={"User-Agent": "JT-Code-RAG/1.0"},
            timeout=settings.RAG_URL_FETCH_TIMEOUT_SECONDS,
            max_bytes=settings.RAG_MAX_EXTRACTED_BYTES,
        )
    except (EgressDenied, EgressError) as exc:
        raise ExtractionError(f"URL fetch failed: {exc}") from exc
    if response.status_code >= 400:
        raise ExtractionError(f"URL fetch returned HTTP {response.status_code}.")
    extension = url.rsplit("?", 1)[0].rsplit(".", 1)[-1] if "." in url.rsplit("/", 1)[-1] else ""
    return _from_bytes(
        response.content, content_type=response.headers.get("content-type", ""), extension=extension
    )


def _extract_pdf(content: bytes) -> ExtractedText:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        pages = [_normalize(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:  # pypdf raises several exception types
        raise ExtractionError(f"PDF extraction failed: {exc}") from exc
    offsets: list[int] = []
    parts: list[str] = []
    cursor = 0
    for page in pages:
        offsets.append(cursor)
        parts.append(page)
        cursor += len(page) + 2
    text = "\n\n".join(parts)
    if not text.strip():
        raise ExtractionError("The PDF has no extractable text (scanned PDFs need OCR).")
    return ExtractedText(text=_cap_size(text), mime_type="application/pdf", page_offsets=offsets)


def _extract_docx(content: bytes) -> ExtractedText:
    try:
        from docx import Document as DocxDocument

        document = DocxDocument(BytesIO(content))
    except Exception as exc:  # python-docx raises several exception types
        raise ExtractionError(f"DOCX extraction failed: {exc}") from exc
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith("heading"):
            level = style.rsplit(" ", 1)[-1]
            blocks.append(f"{'#' * (int(level) if level.isdigit() else 1)} {text}")
        else:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return _result("\n\n".join(blocks), _DOCX)


def _extract_asset_text(asset_id: str, metadata: dict[str, Any]) -> ExtractedText:
    """Read a registered asset through the same egress and size controls as URLs."""
    from urllib.parse import urlsplit

    from apps.assets.imagekit import generate_signed_delivery_url
    from apps.assets.models import Asset

    asset = Asset.objects.filter(id=asset_id, status=Asset.Status.READY).first()
    if asset is None:
        raise ExtractionError("The requested knowledge asset is unavailable.")
    expected_org = metadata.get("organization_id")
    if not expected_org or str(asset.organization_id) != str(expected_org):
        raise ExtractionError("The requested knowledge asset belongs to another organization.")
    if asset.bytes > settings.RAG_MAX_EXTRACTED_BYTES * 4:
        raise ExtractionError("The knowledge asset exceeds the ingestion size limit.")
    host = urlsplit(settings.IMAGEKIT_ENDPOINT_URL).hostname
    if not host:
        raise ExtractionError("ImageKit delivery is not configured.")
    try:
        response = safe_request(
            "GET",
            generate_signed_delivery_url(asset.imagekit_file_path),
            allowed_hosts=[host],
            timeout=settings.RAG_URL_FETCH_TIMEOUT_SECONDS,
            max_bytes=settings.RAG_MAX_EXTRACTED_BYTES * 4,
        )
    except (EgressDenied, EgressError) as exc:
        raise ExtractionError(f"Asset fetch failed: {exc}") from exc
    if response.status_code >= 400:
        raise ExtractionError(f"Asset fetch returned HTTP {response.status_code}.")
    if asset.checksum_sha256:
        import hashlib

        if hashlib.sha256(response.content).hexdigest() != asset.checksum_sha256:
            raise ExtractionError("The knowledge asset failed its integrity check.")
    content_type = str(asset.metadata.get("content_type") or response.headers.get("content-type", ""))
    return _from_bytes(response.content, content_type=content_type, extension=asset.format)
