"""Excel chart sheets for the visualization engine (openpyxl, runs inside the isolated engine).

Each chart becomes one worksheet: its title and takeaway on top, the organized values it shows
as a table, and — when graphs are requested — a native Excel chart drawn from those cells, so
the chart stays editable and updates when the person changes a value.
"""

from __future__ import annotations

from typing import Any

HEADER_FILL = "4F46E5"
TITLE_ROW, SUBTITLE_ROW, NOTE_ROW, HEADER_ROW = 1, 2, 3, 5
MAX_CATEGORIES = 120
MAX_SERIES = 8
MAX_SLICES = 30
NO_GRAPH = {"sankey", "graph", "parallel", "lines", "custom", "gauge"}


def _number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def numeric_columns(columns: list[str], rows: list[list[Any]]) -> list[int]:
    found = []
    for index in range(len(columns)):
        values = [row[index] for row in rows if index < len(row) and row[index] is not None]
        if values and all(_number(value) for value in values):
            found.append(index)
    return found


def series_types(option: dict[str, Any]) -> list[str]:
    return [str(item.get("type")) for item in option.get("series") or [] if isinstance(item, dict)]


def _axis(option: dict[str, Any], key: str) -> dict[str, Any]:
    value = option.get(key)
    if isinstance(value, list):
        return value[0] if value and isinstance(value[0], dict) else {}
    return value if isinstance(value, dict) else {}


def chart_kind(option: dict[str, Any], columns: list[str], rows: list[list[Any]]) -> str:
    """Which Excel chart fits the table this ECharts option produced ('' = values only)."""
    kinds = set(series_types(option))
    if not rows or len(columns) < 2 or kinds <= NO_GRAPH:
        return ""
    if kinds & {"pie", "funnel"}:
        return "pie"
    if "map" in kinds or kinds & {"treemap", "sunburst", "tree"}:
        return "bar_h"
    if "radar" in kinds:
        return "radar"
    if "heatmap" in kinds and not kinds & {"bar", "line"}:
        x_axis, y_axis = _axis(option, "xAxis"), _axis(option, "yAxis")
        return "colour_scale" if x_axis.get("type") == "category" == y_axis.get("type") else "line"
    if "boxplot" in kinds:
        return "dots"
    if "candlestick" in kinds:
        return "line"
    if columns[:1] == ["Series"] and len(columns) >= 3:
        numeric = numeric_columns(columns, rows)
        if 2 not in numeric:
            return ""
        return "scatter" if 1 in numeric else "series_lines"
    if "bar" in kinds and "line" in kinds:
        return "combo"
    if "bar" in kinds:
        return "bar_h" if _axis(option, "yAxis").get("type") == "category" else "bar"
    if "line" in kinds:
        series = [item for item in option.get("series") or [] if item.get("type") == "line"]
        return "area" if series and all(item.get("areaStyle") is not None for item in series) else "line"
    return "bar" if numeric_columns(columns, rows) else ""


def _colour(palette: list[str], index: int) -> str | None:
    if not palette:
        return None
    value = str(palette[index % len(palette)])
    return value[1:].upper() if value.startswith("#") and len(value) == 7 else None


def _paint(series: Any, colour: str | None, line: bool = False) -> None:
    if not colour:
        return
    if line:
        series.graphicalProperties.line.solidFill = colour
        series.graphicalProperties.line.width = 22000
        series.smooth = False
    else:
        series.graphicalProperties.solidFill = colour
        series.graphicalProperties.line.solidFill = colour


def _show_axes(chart: Any) -> None:
    # Newer Excel builds hide axes that openpyxl does not explicitly keep.
    for axis in (getattr(chart, "x_axis", None), getattr(chart, "y_axis", None)):
        if axis is not None:
            axis.delete = False


def write_table(sheet: Any, columns: list[str], rows: list[list[Any]], start_row: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=start_row, column=index, value=str(column))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for offset, row in enumerate(rows, start=1):
        for index, value in enumerate(row[: len(columns)], start=1):
            if isinstance(value, (list, dict)):
                value = str(value)
            cell = sheet.cell(row=start_row + offset, column=index, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # never let a data value become a formula
    for index, column in enumerate(columns, start=1):
        sample = [str(column)] + [str(row[index - 1]) for row in rows[:200] if index - 1 < len(row)]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(text) for text in sample) + 2, 40
        )


def _groups(rows: list[list[Any]], first: int) -> list[tuple[str, int, int]]:
    """Consecutive rows of the same series (rows are sorted by series before writing)."""
    groups: list[tuple[str, int, int]] = []
    for offset, row in enumerate(rows):
        name = str(row[0] or "")
        if groups and groups[-1][0] == name:
            groups[-1] = (name, groups[-1][1], first + offset)
        else:
            groups.append((name, first + offset, first + offset))
    return groups


def _references(sheet: Any, value_cols: list[int], last: int) -> list[Any]:
    from openpyxl.chart import Reference

    return [Reference(sheet, min_col=index + 1, min_row=HEADER_ROW, max_row=last) for index in value_cols]


def build_chart(
    sheet: Any, kind: str, option: dict[str, Any], columns: list[str], rows: list[list[Any]], title: str
) -> Any:
    """A native Excel chart over the table written at HEADER_ROW (None when nothing fits)."""
    from openpyxl.chart import (
        AreaChart,
        BarChart,
        DoughnutChart,
        LineChart,
        PieChart,
        RadarChart,
        Reference,
        ScatterChart,
        Series,
    )
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint

    palette = [str(colour) for colour in option.get("color") or [] if isinstance(colour, str)]
    numeric = numeric_columns(columns, rows)
    first = HEADER_ROW + 1
    x_axis, y_axis = _axis(option, "xAxis"), _axis(option, "yAxis")
    x_name, y_name = str(x_axis.get("name") or ""), str(y_axis.get("name") or "")
    chart: Any

    if kind == "pie":
        value_col = next((index for index in numeric if index > 0), None)
        if value_col is None:
            return None
        last = HEADER_ROW + min(len(rows), MAX_SLICES)
        radius = next((item.get("radius") for item in option.get("series") or [] if item.get("radius")), None)
        hole = isinstance(radius, list) and str(radius[0]).rstrip("%") not in {"0", ""}
        chart = DoughnutChart(holeSize=55) if hole else PieChart()
        chart.add_data(
            Reference(sheet, min_col=value_col + 1, min_row=HEADER_ROW, max_row=last), titles_from_data=True
        )
        chart.set_categories(Reference(sheet, min_col=1, min_row=first, max_row=last))
        series = chart.series[0]
        for index in range(last - HEADER_ROW):
            colour = _colour(palette, index)
            if colour:
                point = DataPoint(idx=index)
                point.graphicalProperties.solidFill = colour
                series.dPt.append(point)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.legend.position = "r"
    elif kind in {"scatter", "series_lines"}:
        x_col, y_col = 1, 2
        groups = _groups(rows, first)
        line_names = {
            str(item.get("name") or "")
            for item in option.get("series") or []
            if isinstance(item, dict) and item.get("type") == "line"
        }
        if kind == "scatter":
            chart = ScatterChart()
            chart.style = 13
        else:
            chart = LineChart()
        for index, (name, start, end) in enumerate(groups[:MAX_SERIES]):
            values = Reference(sheet, min_col=y_col + 1, min_row=start, max_row=end)
            colour = _colour(palette, index)
            if kind == "series_lines":
                series = Series(values, title=name or columns[y_col])
                _paint(series, colour, line=True)
                chart.series.append(series)
                continue
            x_values = Reference(sheet, min_col=x_col + 1, min_row=start, max_row=end)
            series = Series(values, x_values, title=name or columns[y_col])
            if name in line_names:
                _paint(series, colour, line=True)
                series.marker.symbol = "none"
            else:
                series.marker.symbol = "circle"
                series.marker.size = 5
                series.graphicalProperties.line.noFill = True
                if colour:
                    series.marker.graphicalProperties.solidFill = colour
                    series.marker.graphicalProperties.line.solidFill = colour
            chart.series.append(series)
        if kind == "series_lines" and groups:
            _, start, end = max(groups[:MAX_SERIES], key=lambda group: group[2] - group[1])
            chart.set_categories(Reference(sheet, min_col=x_col + 1, min_row=start, max_row=end))
        chart.x_axis.title = x_name or columns[x_col]
        chart.y_axis.title = y_name or columns[y_col]
        if len(groups) < 2:
            chart.legend = None
    elif kind == "radar":
        if len(numeric) < 3:
            return None
        chart = RadarChart()
        chart.type = "marker"
        for index, row in enumerate(range(first, first + min(len(rows), MAX_SERIES))):
            chart.add_data(
                Reference(sheet, min_col=1, max_col=len(columns), min_row=row, max_row=row),
                from_rows=True,
                titles_from_data=True,
            )
            _paint(chart.series[-1], _colour(palette, index), line=True)
        chart.set_categories(
            Reference(sheet, min_col=2, max_col=len(columns), min_row=HEADER_ROW, max_row=HEADER_ROW)
        )
    else:
        value_cols = [index for index in numeric if index > 0][:MAX_SERIES]
        if not value_cols:
            return None
        last = HEADER_ROW + min(len(rows), MAX_CATEGORIES)
        categories = Reference(sheet, min_col=1, min_row=first, max_row=last)
        horizontal = kind == "bar_h"
        stacked = any(item.get("stack") for item in option.get("series") or [] if isinstance(item, dict))
        if kind == "combo":
            types = {
                str(item.get("name") or ""): item.get("type")
                for item in option.get("series") or []
                if item.get("type") not in {"custom", "lines"}
            }
            bar_cols = [index for index in value_cols if types.get(columns[index]) == "bar"] or value_cols[:1]
            line_cols = [index for index in value_cols if index not in bar_cols]
            chart = BarChart()
            for index, ref in enumerate(_references(sheet, bar_cols, last)):
                chart.add_data(ref, titles_from_data=True)
                _paint(chart.series[-1], _colour(palette, index))
            chart.set_categories(categories)
            if line_cols:
                line = LineChart()
                for index, ref in enumerate(_references(sheet, line_cols, last), start=len(bar_cols)):
                    line.add_data(ref, titles_from_data=True)
                    _paint(line.series[-1], _colour(palette, index), line=True)
                line.y_axis.axId = 200
                line.y_axis.crosses = "max"
                line.y_axis.title = columns[line_cols[0]]
                line.y_axis.delete = False
                chart += line
        else:
            if kind in {"bar", "bar_h"}:
                chart = BarChart()
                chart.type = "bar" if horizontal else "col"
                if stacked:
                    chart.grouping = "stacked"
                    chart.overlap = 100
            elif kind == "area":
                chart = AreaChart()
                if stacked:
                    chart.grouping = "stacked"
            else:
                chart = LineChart()
            for index, ref in enumerate(_references(sheet, value_cols, last)):
                chart.add_data(ref, titles_from_data=True)
                series = chart.series[-1]
                if kind in {"line", "dots"}:
                    _paint(series, _colour(palette, index), line=True)
                    if kind == "dots":
                        series.graphicalProperties.line.noFill = True
                        series.marker.symbol = "circle"
                        series.marker.size = 7
                else:
                    _paint(series, _colour(palette, index))
            chart.set_categories(categories)
        chart.x_axis.title = (y_name if horizontal else x_name) or columns[0]
        chart.y_axis.title = (x_name if horizontal else y_name) or (
            columns[value_cols[0]] if len(value_cols) == 1 else None
        )
        if len(chart.series) < 2:
            chart.legend = None
    _show_axes(chart)
    chart.title = title[:250]
    if chart.legend is not None and kind != "pie":
        chart.legend.position = "b"
    chart.width, chart.height = 24, 12
    return chart


def add_chart_sheet(
    book: Any,
    name: str,
    chart: dict[str, Any],
    table: dict[str, Any],
    graphs: bool,
) -> bool:
    """Write one chart's values (and its Excel chart) to a new sheet; True when a graph was drawn."""
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    sheet = book.create_sheet(name)
    option = chart.get("option") or {}
    title = str(chart.get("title") or name)
    columns = [str(column) for column in table.get("columns") or []]
    rows = [list(row) for row in table.get("rows") or []]
    kind = chart_kind(option, columns, rows)
    if kind in {"scatter", "series_lines"}:
        order = {name: index for index, name in enumerate(dict.fromkeys(str(row[0] or "") for row in rows))}
        rows.sort(key=lambda row: order[str(row[0] or "")])
    sheet.cell(row=TITLE_ROW, column=1, value=title).font = Font(bold=True, size=14, color="312E81")
    if chart.get("subtitle"):
        sheet.cell(row=SUBTITLE_ROW, column=1, value=str(chart["subtitle"])).font = Font(
            italic=True, color="64748B"
        )
    note = str(chart.get("note") or "")
    if table.get("truncated"):
        note = f"{note} Showing the first {len(rows)} rows.".strip()
    if note:
        sheet.cell(row=NOTE_ROW, column=1, value=note[:1000]).font = Font(color="334155")
    write_table(sheet, columns, rows, HEADER_ROW)
    sheet.freeze_panes = sheet.cell(row=HEADER_ROW + 1, column=1)
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if not graphs or not kind:
        return False
    if kind == "colour_scale":
        numeric = numeric_columns(columns, rows)
        if numeric and rows:
            area = (
                f"{get_column_letter(min(numeric) + 1)}{HEADER_ROW + 1}:"
                f"{get_column_letter(max(numeric) + 1)}{HEADER_ROW + len(rows)}"
            )
            sheet.conditional_formatting.add(
                area,
                ColorScaleRule(start_type="min", start_color="EEF2FF", end_type="max", end_color="4F46E5"),
            )
            return True
        return False
    excel_chart = build_chart(sheet, kind, option, columns, rows, title)
    if excel_chart is None:
        return False
    sheet.add_chart(excel_chart, f"{get_column_letter(len(columns) + 2)}{HEADER_ROW}")
    return True
