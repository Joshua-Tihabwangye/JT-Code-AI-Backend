"""Isolated visualization engine (pandas, NumPy, SciPy, statsmodels, scikit-learn, openpyxl).

Runs exactly like :mod:`apps.analytics.engine`: as a separate ``python -I``
process with an empty environment, a private working directory and kernel
resource limits applied before any untrusted data is read. It reads one JSON
request from stdin and writes one JSON response to stdout.

Requests never carry code or expressions: they name a chart or prediction task
from the catalog, map its slots to columns and set validated options. Charts are
returned as Apache ECharts option JSON (see :mod:`viz_charts`).
"""

from __future__ import annotations

import base64
import importlib.util
import io
import json
import math
import sys
from pathlib import Path
from typing import Any

_HERE = Path(__file__).resolve().parent


def _load(name: str, filename: str) -> Any:
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, _HERE / filename)
    if spec is None or spec.loader is None:  # pragma: no cover - fixed, packaged paths
        raise ImportError(name)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


d = _load("jt_viz_data", "viz_data.py")
c = _load("jt_viz_charts", "viz_charts.py")
q = _load("jt_viz_quality", "viz_quality.py")
x = _load("jt_viz_excel", "viz_excel.py")
vq = _load("jt_viz_query", "viz_query.py")
vv = _load("jt_viz_visuals", "viz_visuals.py")
base = d.base
catalog = d.catalog
AnalysisError = d.AnalysisError
MODEL_ROWS = d.MODEL_ROWS
GRAIN_FREQ = d.GRAIN_FREQ
SEASON = d.SEASON
POSITIVE, NEGATIVE = c.POSITIVE, c.NEGATIVE
num, nums, compact, fn = c.num, c.nums, c.compact, c.fn

# Re-exported for callers and tests.
column_kinds = d.column_kinds
typed = d.typed
render = c.render
RENDERERS = c.RENDERERS
THEME_COLORS = c.THEME_COLORS


def _chart(
    chart_id: str, title: str, option: dict[str, Any], ctx: dict[str, Any], how: str, insight: str = ""
) -> dict[str, Any]:
    option = c.finish(option, ctx)
    return {
        "id": chart_id,
        "title": title,
        "option": option,
        "table": c.option_table(option),
        "explanation": {"about": title, "howToRead": how, "insight": insight},
    }


def _scatter_xy(
    xs: Any,
    ys: Any,
    ctx: dict[str, Any],
    x_name: str,
    y_name: str,
    *,
    diagonal: bool = False,
    colours: list[str] | None = None,
) -> dict[str, Any]:
    option = c.base(trigger="item", legend=diagonal)
    option["xAxis"] = c.value_axis(x_name, scale=True)
    option["yAxis"] = c.value_axis(y_name, scale=True)
    points = (
        [
            {"value": [num(a), num(b)], "itemStyle": {"color": colours[i]}}
            for i, (a, b) in enumerate(zip(xs, ys, strict=True))
        ]
        if colours
        else [[num(a), num(b)] for a, b in zip(xs, ys, strict=True)]
    )
    option["series"] = [
        {
            "type": "scatter",
            "name": "Rows",
            "data": points,
            "symbolSize": 6,
            "itemStyle": {"opacity": 0.65},
            "large": len(points) > 2000 and not colours,
        }
    ]
    if diagonal:
        low, high = float(min(min(xs), min(ys))), float(max(max(xs), max(ys)))
        option["series"].append(
            {
                "type": "line",
                "name": "Perfect",
                "data": [[num(low), num(low)], [num(high), num(high)]],
                "showSymbol": False,
                "lineStyle": {"type": "dashed", "color": NEGATIVE},
            }
        )
    return option


def _importance(
    model: Any, matrix: Any, x_test: Any, y_test: Any, ctx: dict[str, Any]
) -> tuple[dict[str, Any], list[str]]:
    from sklearn.inspection import permutation_importance

    result = permutation_importance(model, x_test, y_test, n_repeats=4, random_state=0, n_jobs=1)
    order = result.importances_mean.argsort()[::-1][:15][::-1]
    option = c._importance_option(
        [str(matrix.columns[i]) for i in order],
        result.importances_mean[order],
        result.importances_std[order],
        ctx,
    )
    top = [str(matrix.columns[i]) for i in result.importances_mean.argsort()[::-1][:3]]
    chart = _chart(
        "feature_importance",
        "Which inputs matter most",
        option,
        ctx,
        catalog.READ_GUIDE["feature_importance"],
        f"The strongest drivers are {', '.join(top)}.",
    )
    return chart, top


def _pick_model(candidates, chosen, x, y, scoring, ctx):
    from sklearn.model_selection import cross_val_score

    if chosen != "auto":
        return chosen, []
    comparison = []
    if len(x) > 4000:
        sample = x.sample(4000, random_state=0)
        sample_x, sample_y = sample, y.loc[sample.index]
    else:
        sample_x, sample_y = x, y
    for name, factory in candidates:
        try:
            score = float(
                cross_val_score(factory(), sample_x, sample_y, cv=3, scoring=scoring, n_jobs=1).mean()
            )
        except ValueError:
            continue
        comparison.append({"model": name, "score": round(score, 6)})
    if not comparison:
        raise AnalysisError("No model could be trained on this data.")
    best = max(comparison, key=lambda item: item["score"])["model"]
    ctx["notes"].append(f"Automatic selection compared {len(comparison)} models by 3-fold cross-validation.")
    return best, comparison


def predict_regression(frame, kinds, fields, options, ctx, model_name):
    import numpy as np
    from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
    from sklearn.model_selection import train_test_split

    target = fields["target"]
    data = frame.assign(**{target: d._numeric(frame, target)}).dropna(subset=[target])
    data = d._sample(data, MODEL_ROWS, ctx, "model")
    if len(data.index) < 20:
        raise AnalysisError("At least 20 rows with a target value are needed.")
    features = d._inputs(data, fields, kinds, target, ctx)
    matrix = d._design_matrix(data, features, kinds, ctx)
    y = data[target]
    names = ("linear", "ridge", "random_forest", "gradient_boosting")
    model_name, comparison = _pick_model(
        [(n, lambda n=n: d._estimator(n, True)) for n in names], model_name, matrix, y, "r2", ctx
    )
    x_train, x_test, y_train, y_test = train_test_split(
        matrix, y, test_size=float(options.get("test_size", "0.2")), random_state=0
    )
    model = d._estimator(model_name, True).fit(x_train, y_train)
    predicted = model.predict(x_test)
    r2 = float(r2_score(y_test, predicted))
    mae = float(mean_absolute_error(y_test, predicted))
    rmse = float(math.sqrt(mean_squared_error(y_test, predicted)))
    baseline = float(mean_absolute_error(y_test, np.full(len(y_test), y_train.mean())))
    nonzero = (y_test != 0).to_numpy()
    mape = (
        float(
            np.mean(np.abs((y_test.to_numpy()[nonzero] - predicted[nonzero]) / y_test.to_numpy()[nonzero]))
            * 100
        )
        if nonzero.any()
        else float("nan")
    )
    metrics = [
        d._metric(
            "r2",
            "R² (hold-out)",
            r2,
            "number",
            d._quality("r2", r2),
            "Share of the variation the model explains (1 is perfect).",
        ),
        d._metric("mae", "Mean absolute error", mae, "number", None, f"Average miss in {target} units."),
        d._metric("rmse", "Root mean squared error", rmse, "number", None, "Penalises large misses more."),
        d._metric(
            "mape",
            "Mean absolute % error",
            mape,
            "percent_value",
            d._quality("mape", mape) if not math.isnan(mape) else None,
        ),
        d._metric(
            "baseline_mae",
            "Baseline MAE (predict the mean)",
            baseline,
            "number",
            None,
            "The model should beat this.",
        ),
        d._metric("rows", "Training rows", int(len(x_train)), "integer"),
    ]
    residual = y_test.to_numpy() - predicted
    residual_option = _scatter_xy(
        predicted,
        residual,
        ctx,
        "Predicted",
        "Residual (actual − predicted)",
        colours=[POSITIVE if r >= 0 else NEGATIVE for r in residual],
    )
    residual_option["series"][0]["markLine"] = {
        "silent": True,
        "symbol": "none",
        "data": [{"yAxis": 0}],
        "lineStyle": {"type": "dashed", "color": ctx["tokens"]["muted"]},
    }
    importance, top = _importance(model, matrix, x_test, y_test, ctx)
    numeric_inputs = [column for column in matrix.columns if kinds.get(column) == "numeric"]
    baseline_mean = float(np.mean(predicted)) or 1.0
    sensitivity = []
    for column in [name for name in top if name in numeric_inputs] + [
        n for n in numeric_inputs if n not in top
    ]:
        if len(sensitivity) >= 6:
            break
        effects = []
        for factor in (0.9, 1.1):
            shifted = x_test.copy()
            shifted[column] = shifted[column] * factor
            effects.append((float(np.mean(model.predict(shifted))) - baseline_mean) / abs(baseline_mean))
        sensitivity.append((column, effects[0], effects[1]))
    what_if_chart = None
    if sensitivity:
        sensitivity.sort(key=lambda item: abs(item[2]) + abs(item[1]))
        tornado = c.base()
        tornado["tooltip"]["axisPointer"] = {"type": "shadow"}
        tornado["tooltip"]["valueFormatter"] = fn("percent", digits=1, scale=100)
        tornado["xAxis"] = {
            "type": "value",
            "name": f"Change in predicted {target}",
            "nameLocation": "middle",
            "nameGap": 28,
            "axisLabel": {"formatter": fn("percent", digits=0, scale=100)},
        }
        tornado["yAxis"] = c.category_axis([item[0] for item in sensitivity])
        tornado["series"] = [
            {
                "type": "bar",
                "name": "Input −10%",
                "stack": "w",
                "data": [num(item[1], 4) for item in sensitivity],
                "itemStyle": {"color": NEGATIVE},
            },
            {
                "type": "bar",
                "name": "Input +10%",
                "stack": "w",
                "data": [num(item[2], 4) for item in sensitivity],
                "itemStyle": {"color": POSITIVE},
            },
        ]
        strongest = max(sensitivity, key=lambda item: abs(item[2]))
        what_if_chart = _chart(
            "what_if",
            "What if an input changes by 10%?",
            tornado,
            ctx,
            "Each row raises (green) or lowers (red) one input by 10% and shows how the predicted "
            "outcome moves.",
            f"Raising {strongest[0]} by 10% changes predicted {target} by {strongest[2]:+.1%}.",
        )
    charts = [
        _chart(
            "predicted_vs_actual",
            "Predicted vs actual (hold-out)",
            _scatter_xy(
                y_test.to_numpy(), predicted, ctx, f"Actual {target}", f"Predicted {target}", diagonal=True
            ),
            ctx,
            catalog.READ_GUIDE["predicted_vs_actual"],
            f"The model explains {max(r2, 0):.0%} of the variation on unseen rows.",
        ),
        _chart(
            "residual_plot",
            "Residuals",
            residual_option,
            ctx,
            catalog.READ_GUIDE["residual_plot"],
            f"Typical error ±{compact(mae)}.",
        ),
        importance,
    ]
    if what_if_chart:
        charts.append(what_if_chart)
    result = data.copy()
    result[f"predicted_{target}"] = model.predict(matrix)
    result["split"] = np.where(result.index.isin(x_test.index), "test", "train")
    import pandas as pd

    table = pd.DataFrame({"actual": y_test, "predicted": predicted, "error": residual}).head(200)
    summary = [
        (
            f"The {d._model_names()[model_name]} explains {max(r2, 0):.0%} of the variation in {target} "
            "on rows it did not see."
        ),
        (
            f"Typical error is ±{mae:,.3g} ({'better' if mae < baseline else 'no better'} than always "
            f"predicting the average, ±{baseline:,.3g})."
        ),
        f"The strongest drivers are {', '.join(top)}.",
    ]
    return metrics, charts, table, result, summary, comparison, model_name


def predict_classification(frame, kinds, fields, options, ctx, model_name):
    import numpy as np
    import pandas as pd
    from sklearn.metrics import (
        accuracy_score,
        balanced_accuracy_score,
        confusion_matrix,
        f1_score,
        precision_recall_curve,
        roc_auc_score,
        roc_curve,
    )
    from sklearn.model_selection import train_test_split

    target = fields["target"]
    data = frame.dropna(subset=[target])
    data = d._sample(data, MODEL_ROWS, ctx, "model")
    labels = d._label(data[target])
    counts = labels.value_counts()
    if len(counts) < 2:
        raise AnalysisError("The target needs at least two categories.")
    if len(counts) > 30:
        raise AnalysisError("The target has more than 30 categories; choose a column with fewer classes.")
    keep = labels.isin(counts[counts >= 2].index)
    data, labels = data[keep], labels[keep]
    features = d._inputs(data, fields, kinds, target, ctx)
    matrix = d._design_matrix(data, features, kinds, ctx)
    names = ("logistic", "random_forest", "gradient_boosting")
    model_name, comparison = _pick_model(
        [(n, lambda n=n: d._estimator(n, False)) for n in names], model_name, matrix, labels, "f1_macro", ctx
    )
    x_train, x_test, y_train, y_test = train_test_split(
        matrix, labels, test_size=float(options.get("test_size", "0.2")), random_state=0, stratify=labels
    )
    model = d._estimator(model_name, False).fit(x_train, y_train)
    predicted = model.predict(x_test)
    probabilities = model.predict_proba(x_test)
    classes = [str(value) for value in model.classes_]
    accuracy = float(accuracy_score(y_test, predicted))
    baseline = float((y_test == y_train.value_counts().idxmax()).mean())
    f1 = float(f1_score(y_test, predicted, average="macro", zero_division=0))
    try:
        auc = float(
            roc_auc_score(y_test, probabilities[:, 1])
            if len(classes) == 2
            else roc_auc_score(y_test, probabilities, multi_class="ovr", labels=classes)
        )
    except ValueError:
        auc = float("nan")
    metrics = [
        d._metric(
            "accuracy",
            "Accuracy (hold-out)",
            accuracy,
            "percent",
            d._quality("accuracy", accuracy, baseline),
            "Share of rows classified correctly.",
        ),
        d._metric(
            "baseline",
            "Baseline accuracy (most common class)",
            baseline,
            "percent",
            None,
            "The model should beat this.",
        ),
        d._metric(
            "balanced_accuracy",
            "Balanced accuracy",
            float(balanced_accuracy_score(y_test, predicted)),
            "percent",
        ),
        d._metric("f1", "F1 (macro)", f1, "number", d._quality("f1", f1)),
        d._metric(
            "auc",
            "ROC AUC",
            auc,
            "number",
            d._quality("auc", auc) if not math.isnan(auc) else None,
            "Ranking quality: 0.5 is random, 1 is perfect.",
        ),
        d._metric("rows", "Training rows", int(len(x_train)), "integer"),
    ]
    matrix_counts = confusion_matrix(y_test, predicted, labels=classes)
    cells = [[j, i, int(matrix_counts[i, j])] for i in range(len(classes)) for j in range(len(classes))]
    confusion = c._heatmap(classes, classes, cells, ctx, name="Rows")
    confusion["xAxis"]["name"], confusion["yAxis"]["name"] = "Predicted", "Actual"
    roc = c.base()
    c._unit_axes(roc, "False positive rate", "True positive rate")
    for label in classes[:8] if len(classes) > 2 else classes[1:]:
        truth = (y_test == label).astype(int)
        if truth.nunique() < 2:
            continue
        fpr, tpr, _ = roc_curve(truth, probabilities[:, classes.index(label)])
        roc["series"].append(
            {
                "type": "line",
                "name": label if len(classes) > 2 else f"ROC (AUC {auc:.3f})",
                "data": c._thin(fpr, tpr),
                "showSymbol": False,
                "lineStyle": {"width": 3},
            }
        )
    roc["series"].append(
        {
            "type": "line",
            "name": "Random",
            "data": [[0, 0], [1, 1]],
            "showSymbol": False,
            "lineStyle": {"type": "dashed", "color": ctx["tokens"]["muted"]},
        }
    )
    importance, top = _importance(model, matrix, x_test, y_test, ctx)
    charts = [
        _chart(
            "confusion_matrix",
            "Confusion matrix (hold-out)",
            confusion,
            ctx,
            catalog.READ_GUIDE["confusion_matrix"],
            f"{accuracy:.0%} of unseen rows are classified correctly.",
        ),
        _chart(
            "roc_curve",
            "ROC curve" + (" (one vs rest)" if len(classes) > 2 else ""),
            roc,
            ctx,
            catalog.READ_GUIDE["roc_curve"],
            "" if math.isnan(auc) else f"AUC = {auc:.3f}.",
        ),
        importance,
    ]
    if len(classes) == 2:
        precision, recall, _ = precision_recall_curve((y_test == classes[1]).astype(int), probabilities[:, 1])
        pr = c.base()
        c._unit_axes(pr, "Recall", "Precision")
        pr["series"] = [
            {
                "type": "line",
                "name": f"Positive: {classes[1]}",
                "data": c._thin(recall, precision),
                "showSymbol": False,
                "areaStyle": {"opacity": 0.15},
                "lineStyle": {"width": 3},
            }
        ]
        charts.insert(
            2, _chart("pr_curve", "Precision–recall curve", pr, ctx, catalog.READ_GUIDE["pr_curve"])
        )
    all_probabilities = model.predict_proba(matrix)
    result = data.copy()
    result[f"predicted_{target}"] = model.predict(matrix)
    result["confidence"] = all_probabilities.max(axis=1).round(4)
    for index, label in enumerate(classes[:10]):
        result[f"p_{label}"] = all_probabilities[:, index].round(4)
    result["split"] = np.where(result.index.isin(x_test.index), "test", "train")
    table = pd.DataFrame(
        {"actual": y_test, "predicted": predicted, "confidence": probabilities.max(axis=1).round(4)}
    ).head(200)
    summary = [
        (
            f"The {d._model_names()[model_name]} classifies {accuracy:.0%} of unseen rows correctly "
            f"(guessing the most common class gives {baseline:.0%})."
        ),
        f"The strongest drivers are {', '.join(top)}.",
    ]
    return metrics, charts, table, result, summary, comparison, model_name


def predict_clustering(frame, kinds, fields, options, ctx, model_name):
    import numpy as np
    import pandas as pd
    from sklearn.cluster import AgglomerativeClustering, KMeans
    from sklearn.decomposition import PCA
    from sklearn.metrics import calinski_harabasz_score, davies_bouldin_score, silhouette_score
    from sklearn.mixture import GaussianMixture
    from sklearn.preprocessing import StandardScaler

    data = d._sample(frame, MODEL_ROWS, ctx, "model")
    features = d._inputs(data, fields, kinds, "", ctx)
    matrix = d._design_matrix(data, features, kinds, ctx)
    values = StandardScaler().fit_transform(matrix)
    if len(values) < 10:
        raise AnalysisError("At least 10 rows are needed to find segments.")
    k = int(options.get("clusters") or 0)
    by_k: list[tuple[int, float]] = []
    if not k:
        k, by_k = d._choose_k(values)
    if model_name == "gaussian_mixture":
        assigned = GaussianMixture(n_components=k, random_state=0).fit_predict(values)
    elif model_name == "agglomerative":
        if len(values) > 5000:
            index = np.random.default_rng(0).choice(len(values), 5000, replace=False)
            sample_labels = AgglomerativeClustering(n_clusters=k).fit_predict(values[index])
            centres = np.array([values[index][sample_labels == cluster].mean(axis=0) for cluster in range(k)])
            assigned = np.argmin(((values[:, None, :] - centres[None]) ** 2).sum(axis=2), axis=1)
            ctx["notes"].append(
                "Hierarchical clustering ran on a 5,000-row sample; other rows joined the nearest segment."
            )
        else:
            assigned = AgglomerativeClustering(n_clusters=k).fit_predict(values)
    else:
        model_name = "kmeans"
        assigned = KMeans(n_clusters=k, n_init=10, random_state=0).fit_predict(values)
    sample = np.random.default_rng(0).choice(len(values), min(len(values), 5000), replace=False)
    silhouette = (
        float(silhouette_score(values[sample], assigned[sample])) if len(set(assigned)) > 1 else float("nan")
    )
    names = pd.Series([f"Segment {label + 1}" for label in assigned])
    metrics = [
        d._metric("clusters", "Segments", int(k), "integer"),
        d._metric(
            "silhouette",
            "Silhouette score",
            silhouette,
            "number",
            d._quality("silhouette", silhouette),
            "Separation of segments: −1 to 1, higher is better.",
        ),
        d._metric(
            "davies_bouldin",
            "Davies–Bouldin index",
            float(davies_bouldin_score(values, assigned)),
            "number",
            None,
            "Lower is better.",
        ),
        d._metric(
            "calinski_harabasz",
            "Calinski–Harabasz score",
            float(calinski_harabasz_score(values, assigned)),
            "number",
            None,
            "Higher is better.",
        ),
    ]
    pca = PCA(n_components=2, random_state=0).fit(values)
    points = pca.transform(values)
    keep = np.random.default_rng(0).choice(len(points), min(len(points), 5000), replace=False)
    order = [f"Segment {i + 1}" for i in range(k)]
    segment_map = c._cluster_scatter(
        points[keep], names.iloc[keep].reset_index(drop=True), (d._axis(pca, 0), d._axis(pca, 1)), ctx
    )
    profile = matrix.assign(Segment=list(names)).groupby("Segment").mean().reindex(order)
    standardized = ((profile - matrix.mean()) / matrix.std().replace(0, 1)).iloc[:, :30]
    cells = [
        [j, i, num(standardized.iat[i, j], 2)]
        for i in range(standardized.shape[0])
        for j in range(standardized.shape[1])
    ]
    heat = c._heatmap(
        [str(col) for col in standardized.columns],
        order,
        cells,
        ctx,
        diverging=True,
        minimum=-2,
        maximum=2,
        name="vs average (σ)",
    )
    sizes = names.value_counts().reindex(order)
    bars = c.base(legend=False)
    bars["xAxis"] = c.category_axis(order)
    bars["yAxis"] = c.value_axis("Rows")
    bars["series"] = [
        {
            "type": "bar",
            "data": [
                {
                    "value": int(v),
                    "itemStyle": {
                        "color": ctx["colors"][i % len(ctx["colors"])],
                        "borderRadius": [4, 4, 0, 0],
                    },
                }
                for i, v in enumerate(sizes.values)
            ],
        }
    ]
    charts = [
        _chart(
            "cluster_map", "Segments on a PCA map", segment_map, ctx, catalog.READ_GUIDE["kmeans_clusters"]
        ),
        _chart(
            "cluster_profile",
            "What makes each segment different",
            heat,
            ctx,
            (
                "Red cells are above the overall average for that input, blue below; read each row as "
                "a segment profile."
            ),
        ),
        _chart(
            "cluster_sizes",
            "Rows per segment",
            bars,
            ctx,
            "Bar height is the number of rows in each segment.",
        ),
    ]
    if by_k:
        silhouette_bars = c.base(legend=False)
        silhouette_bars["xAxis"] = c.category_axis([a for a, _ in by_k], "k")
        silhouette_bars["yAxis"] = c.value_axis("Silhouette")
        silhouette_bars["series"] = [
            {
                "type": "bar",
                "data": [
                    {"value": num(b, 4), "itemStyle": {"color": POSITIVE if a == k else ctx["colors"][0]}}
                    for a, b in by_k
                ],
            }
        ]
        charts.append(
            _chart(
                "silhouette",
                "Choosing the number of segments",
                silhouette_bars,
                ctx,
                catalog.READ_GUIDE["silhouette"],
            )
        )
    result = data.copy()
    result["segment"] = list(names)
    table = profile.round(4).reset_index()
    summary = [
        (
            f"Found {k} segments; the largest is {sizes.idxmax()} with {int(sizes.max()):,} rows "
            f"({sizes.max() / len(names):.0%})."
        ),
        f"Segment separation (silhouette) is {silhouette:.2f}, rated {d._quality('silhouette', silhouette)}.",
    ]
    return (
        metrics,
        charts,
        table,
        result,
        summary,
        [{"model": a, "score": round(b, 6)} for a, b in by_k],
        model_name,
    )


def predict_anomaly(frame, kinds, fields, options, ctx, model_name):
    import numpy as np
    from sklearn.decomposition import PCA
    from sklearn.preprocessing import StandardScaler

    data = d._sample(frame, MODEL_ROWS, ctx, "model")
    features = d._inputs(data, fields, kinds, "", ctx)
    matrix = d._design_matrix(data, features, kinds, ctx)
    values = StandardScaler().fit_transform(matrix)
    contamination = float(options.get("contamination", "0.05"))
    if model_name == "zscore":
        median = np.median(values, axis=0)
        mad = np.median(np.abs(values - median), axis=0) * 1.4826
        scores = (np.abs(values - median) / np.where(mad == 0, 1, mad)).max(axis=1)
        flags = scores > max(float(np.quantile(scores, 1 - contamination)), 3.5)
    elif model_name == "lof":
        from sklearn.neighbors import LocalOutlierFactor

        detector = LocalOutlierFactor(n_neighbors=min(20, len(values) - 1), contamination=contamination)
        flags = detector.fit_predict(values) == -1
        scores = -detector.negative_outlier_factor_
    else:
        from sklearn.ensemble import IsolationForest

        model_name = "isolation_forest"
        forest = IsolationForest(contamination=contamination, random_state=0, n_estimators=200, n_jobs=1).fit(
            values
        )
        flags = forest.predict(values) == -1
        scores = -forest.score_samples(values)
    count = int(flags.sum())
    metrics = [
        d._metric("anomalies", "Anomalies found", count, "integer"),
        d._metric("share", "Share of rows", count / len(values), "percent"),
        d._metric("rows", "Rows checked", int(len(values)), "integer"),
    ]
    points = PCA(n_components=2, random_state=0).fit_transform(values)
    scatter = c.base(trigger="item")
    scatter["xAxis"] = c.value_axis("PC1", scale=True)
    scatter["yAxis"] = c.value_axis("PC2", scale=True)
    scatter["series"] = [
        {
            "type": "scatter",
            "name": "Normal",
            "data": [[num(a), num(b)] for (a, b), bad in zip(points, flags, strict=True) if not bad][:8000],
            "symbolSize": 5,
            "itemStyle": {"color": c.rgba(ctx["colors"][0], 0.45)},
        },
        {
            "type": "scatter",
            "name": "Anomaly",
            "data": [[num(a), num(b)] for (a, b), bad in zip(points, flags, strict=True) if bad],
            "symbolSize": 9,
            "itemStyle": {"color": NEGATIVE},
        },
    ]
    edges = c._bins(scores, 40)
    normal, _ = np.histogram(scores[~flags], bins=edges)
    unusual, _ = np.histogram(scores[flags], bins=edges)
    histogram = c.base()
    histogram["xAxis"] = c.category_axis(
        c._bin_labels(edges),
        "Anomaly score (higher is more unusual)",
        axisLabel={"rotate": 35, "hideOverlap": True},
    )
    histogram["yAxis"] = c.value_axis("Rows")
    histogram["series"] = [
        {
            "type": "bar",
            "name": "Normal",
            "stack": "s",
            "data": [int(v) for v in normal],
            "itemStyle": {"color": c.rgba(ctx["colors"][0], 0.7)},
        },
        {
            "type": "bar",
            "name": "Anomaly",
            "stack": "s",
            "data": [int(v) for v in unusual],
            "itemStyle": {"color": NEGATIVE},
        },
    ]
    charts = [
        _chart(
            "anomaly_map",
            "Unusual rows on a PCA map",
            scatter,
            ctx,
            "Red points are the rows the model flags as unusual compared with the rest.",
        ),
        _chart(
            "anomaly_scores",
            "Anomaly scores",
            histogram,
            ctx,
            "Rows to the right score as more unusual; red bars are the flagged rows.",
        ),
    ]
    result = data.copy()
    result["anomaly_score"] = np.round(scores, 6)
    result["is_anomaly"] = flags
    table = result.sort_values("anomaly_score", ascending=False).head(100)
    summary = [
        f"{count:,} of {len(values):,} rows ({count / len(values):.1%}) look unusual.",
        "The table lists the most unusual rows first.",
    ]
    return metrics, charts, table, result, summary, [], model_name


def _resample(frame, date, target, options):
    """A regular time series (period index) for one measure."""
    series_frame, _, grain = d._time_aggregate(
        frame, date, [target], [], options.get("aggregation", "sum"), options.get("time_grain", "auto")
    )
    if grain not in GRAIN_FREQ:
        diffs = series_frame[date].diff().dropna()
        days = diffs.dt.days.median() if not diffs.empty else 1
        grain = (
            "D" if days <= 1.5 else "W" if days <= 8 else "M" if days <= 32 else "Q" if days <= 95 else "Y"
        )
    freq = GRAIN_FREQ[grain]
    indexed = series_frame.set_index(date)[target]
    if options.get("aggregation", "sum") in {"sum", "count"}:
        series = indexed.resample(freq).sum(min_count=1)
    else:
        series = indexed.resample(freq).mean()
    return series.interpolate(limit_direction="both"), grain, freq


def _forecast_once(name, history, steps, seasonal):
    import numpy as np

    values = history.to_numpy(dtype=float)
    if name == "holt_winters":
        from statsmodels.tsa.holtwinters import ExponentialSmoothing

        fitted = ExponentialSmoothing(
            values,
            trend="add",
            damped_trend=True,
            seasonal="add" if seasonal else None,
            seasonal_periods=seasonal,
            initialization_method="estimated",
        ).fit()
        mean = fitted.forecast(steps)
        spread = float(np.std(values - fitted.fittedvalues)) * np.sqrt(np.arange(1, steps + 1))
        return (
            mean,
            mean - 1.2816 * spread,
            mean + 1.2816 * spread,
            mean - 1.96 * spread,
            mean + 1.96 * spread,
        )
    if name == "arima":
        from statsmodels.tsa.statespace.sarimax import SARIMAX

        seasonal_order = (1, 0, 1, seasonal) if seasonal and len(values) >= 3 * seasonal else (0, 0, 0, 0)
        fitted = SARIMAX(
            values,
            order=(1, 1, 1),
            seasonal_order=seasonal_order,
            enforce_stationarity=False,
            enforce_invertibility=False,
        ).fit(disp=False, maxiter=80)
        result = fitted.get_forecast(steps)
        c80, c95 = result.conf_int(alpha=0.2), result.conf_int(alpha=0.05)
        return result.predicted_mean, c80[:, 0], c80[:, 1], c95[:, 0], c95[:, 1]
    import statsmodels.api as sm

    def design(index):
        columns = [np.ones(len(index)), index]
        if seasonal:
            for harmonic in (1, 2):
                columns += [
                    np.sin(2 * np.pi * harmonic * index / seasonal),
                    np.cos(2 * np.pi * harmonic * index / seasonal),
                ]
        return np.column_stack(columns)

    fitted = sm.OLS(values, design(np.arange(len(values)))).fit()
    prediction = fitted.get_prediction(design(np.arange(len(values), len(values) + steps)))
    p80, p95 = prediction.summary_frame(alpha=0.2), prediction.summary_frame(alpha=0.05)
    return (
        p80["mean"].to_numpy(),
        p80["obs_ci_lower"].to_numpy(),
        p80["obs_ci_upper"].to_numpy(),
        p95["obs_ci_lower"].to_numpy(),
        p95["obs_ci_upper"].to_numpy(),
    )


def _fit_forecast(series, grain, freq, horizon, model_name):
    """Backtest the candidate models, pick the best, and forecast ``horizon`` periods."""
    import numpy as np
    import pandas as pd

    if len(series) < 8:
        raise AnalysisError(f"Forecasting needs at least 8 periods; found {len(series)}.")
    period = SEASON.get(grain)
    seasonal = period if period and len(series) >= 2 * period + 2 else None
    holdout = max(2, min(horizon, len(series) // 5))
    train, test = series.iloc[:-holdout], series.iloc[-holdout:]
    candidates = ["holt_winters", "arima", "linear_trend"] if model_name == "auto" else [model_name]
    comparison, backtests = [], {}
    for name in candidates:
        try:
            mean = _forecast_once(name, train, holdout, seasonal)[0]
        except (ValueError, np.linalg.LinAlgError) as exc:
            if model_name != "auto":
                raise AnalysisError(
                    f"The {name.replace('_', ' ')} model could not be fitted: {str(exc)[:160]}"
                ) from exc
            continue
        mae = float(np.mean(np.abs(test.to_numpy() - mean)))
        comparison.append({"model": name, "score": round(-mae, 6)})
        backtests[name] = (mae, mean)
    if not backtests:
        raise AnalysisError("No forecasting model could be fitted to this series.")
    best = min(backtests, key=lambda name: backtests[name][0])
    mae, backtest_mean = backtests[best]
    actual = test.to_numpy()
    nonzero = actual != 0
    mape = (
        float(np.mean(np.abs((actual[nonzero] - backtest_mean[nonzero]) / actual[nonzero])) * 100)
        if nonzero.any()
        else float("nan")
    )
    rmse = float(np.sqrt(np.mean((actual - backtest_mean) ** 2)))
    mean, lo80, hi80, lo95, hi95 = (
        np.asarray(part, dtype=float) for part in _forecast_once(best, series, horizon, seasonal)
    )
    future = pd.date_range(series.index[-1], periods=horizon + 1, freq=freq)[1:]
    return {
        "best": best,
        "mae": mae,
        "mape": mape,
        "rmse": rmse,
        "holdout": holdout,
        "test": test,
        "backtest": backtest_mean,
        "comparison": comparison,
        "mean": mean,
        "lo80": lo80,
        "hi80": hi80,
        "lo95": lo95,
        "hi95": hi95,
        "future": future,
        "seasonal": seasonal,
    }


def _forecast_option(series, fit, ctx, date, target):
    history_dates, future_dates = c.dates(series.index.to_series()), c.dates(fit["future"].to_series())
    test_dates = c.dates(fit["test"].index.to_series())
    accent = ctx["colors"][2]

    def band(lower, upper, name, alpha, stack):
        return [
            {
                "type": "line",
                "data": [[s, num(v)] for s, v in zip(future_dates, lower, strict=True)],
                "stack": stack,
                "showSymbol": False,
                "lineStyle": {"opacity": 0},
                "tooltip": {"show": False},
                "silent": True,
            },
            {
                "type": "line",
                "name": name,
                "data": [[s, num(h - lo)] for s, h, lo in zip(future_dates, upper, lower, strict=True)],
                "stack": stack,
                "showSymbol": False,
                "lineStyle": {"opacity": 0},
                "areaStyle": {"color": c.rgba(accent, alpha)},
                "tooltip": {"show": False},
            },
        ]

    option = c.base(zoom=len(series) > 60)
    option["xAxis"] = c.time_axis(date)
    option["yAxis"] = c.value_axis(target, scale=True)
    option["series"] = [
        {
            "type": "line",
            "name": "History",
            "data": [[s, num(v)] for s, v in zip(history_dates, series.values, strict=True)],
            "showSymbol": False,
            "lineStyle": {"width": 2.4},
            "itemStyle": {"color": ctx["colors"][0]},
            "markArea": {
                "silent": True,
                "itemStyle": {"color": c.rgba(ctx["tokens"]["muted"], 0.08)},
                "data": [[{"name": "Backtest", "xAxis": test_dates[0]}, {"xAxis": test_dates[-1]}]],
            },
        },
        *band(fit["lo95"], fit["hi95"], "95% range", 0.12, "b95"),
        *band(fit["lo80"], fit["hi80"], "80% range", 0.25, "b80"),
        {
            "type": "line",
            "name": "Forecast",
            "data": [[s, num(v)] for s, v in zip(future_dates, fit["mean"], strict=True)],
            "symbolSize": 6,
            "lineStyle": {"width": 3, "type": "dashed"},
            "itemStyle": {"color": accent},
        },
        {
            "type": "line",
            "name": "Backtest",
            "data": [[s, num(v)] for s, v in zip(test_dates, fit["backtest"], strict=True)],
            "showSymbol": False,
            "lineStyle": {"width": 1.6, "type": "dotted", "color": ctx["tokens"]["muted"]},
        },
    ]
    option["legend"]["data"] = ["History", "Forecast", "80% range", "95% range", "Backtest"]
    return option


def predict_forecast(frame, kinds, fields, options, ctx, model_name):
    import numpy as np
    import pandas as pd

    date, target, split = fields["date"], fields["target"], fields.get("split_by")
    horizon = int(options.get("horizon", 12))
    series, grain, freq = _resample(frame, date, target, options)
    if len(series) > 3000:
        series = series.iloc[-3000:]
        ctx["notes"].append("Using the latest 3,000 periods.")
    fit = _fit_forecast(series, grain, freq, horizon, model_name)
    mean, best = fit["mean"], fit["best"]
    expected, optimistic, pessimistic = (
        float(np.sum(mean)),
        float(np.sum(fit["hi80"])),
        float(np.sum(fit["lo80"])),
    )
    last_period_total = float(series.iloc[-horizon:].sum()) if len(series) >= horizon else float("nan")
    growth = (expected - last_period_total) / abs(last_period_total) if last_period_total else float("nan")
    metrics = [
        d._metric("next", "Next period forecast", float(mean[0]), "number"),
        d._metric("total", f"Expected total (next {horizon})", expected, "number"),
        d._metric("growth", f"Growth vs last {horizon} periods", growth, "percent"),
        d._metric(
            "optimistic",
            "Best case total (80% upper)",
            optimistic,
            "number",
            None,
            "A realistic upside scenario.",
        ),
        d._metric(
            "pessimistic",
            "Worst case total (80% lower)",
            pessimistic,
            "number",
            None,
            "A realistic downside scenario.",
        ),
        d._metric(
            "mape",
            "Backtest error (MAPE)",
            fit["mape"],
            "percent_value",
            d._quality("mape", fit["mape"]) if not math.isnan(fit["mape"]) else None,
            f"Average % miss on the last {fit['holdout']} known periods.",
        ),
        d._metric("mae", "Backtest MAE", fit["mae"], "number"),
        d._metric(
            "seasonality",
            "Seasonality used",
            "yes" if fit["seasonal"] else "no",
            "number",
            None,
            f"Season length {fit['seasonal'] or '—'} periods.",
        ),
    ]
    last = float(series.iloc[-1])
    trend = (mean[-1] - last) / abs(last) if last else float("nan")
    direction = (
        ""
        if math.isnan(trend)
        else f"{target} is expected to {'rise' if trend > 0 else 'fall'} about {abs(trend):.0%} over "
        f"the next {horizon} periods."
    )
    charts = [
        _chart(
            "forecast",
            f"{target}: {horizon}-period forecast",
            _forecast_option(series, fit, ctx, date, target),
            ctx,
            "Solid line: history. Dashed: forecast. Bands: 80% and 95% ranges. Grey area: backtest used "
            "to score the model.",
            direction,
        )
    ]
    scenarios = c.base(legend=False)
    scenarios["tooltip"]["axisPointer"] = {"type": "shadow"}
    scenarios["xAxis"] = c.category_axis(["Worst case", "Expected", "Best case"])
    scenarios["yAxis"] = c.value_axis(f"Total {target}, next {horizon} periods")
    scenarios["series"] = [
        {
            "type": "bar",
            "name": "Total",
            "barMaxWidth": 70,
            "label": {"show": True, "position": "top", "formatter": fn("labelValue")},
            "data": [
                {"value": num(v), "itemStyle": {"color": colour, "borderRadius": [6, 6, 0, 0]}}
                for v, colour in (
                    (pessimistic, NEGATIVE),
                    (expected, ctx["colors"][0]),
                    (optimistic, POSITIVE),
                )
            ],
        }
    ]
    charts.append(
        _chart(
            "scenarios",
            "Best, expected and worst case",
            scenarios,
            ctx,
            "Plan for the expected total; the best and worst cases bracket 80% of likely outcomes.",
            f"Expect about {compact(expected)}, between {compact(pessimistic)} and {compact(optimistic)}.",
        )
    )
    table = pd.DataFrame(
        {
            "period": fit["future"],
            "forecast": mean,
            "lower_80": fit["lo80"],
            "upper_80": fit["hi80"],
            "lower_95": fit["lo95"],
            "upper_95": fit["hi95"],
        }
    ).round(4)
    result = pd.concat(
        [pd.DataFrame({"period": series.index, "actual": series.values}), table], ignore_index=True
    )
    comparison = fit["comparison"]
    if split:
        data = d._labels(frame, [split])
        totals = (
            data.assign(_v=d._numeric(data, target)).groupby(split)["_v"].sum().sort_values(ascending=False)
        )
        groups = list(totals.index[:6])
        lines = c.base(zoom=False)
        lines["xAxis"] = c.time_axis(date)
        lines["yAxis"] = c.value_axis(target, scale=True)
        growth_rows, tables = [], []
        for index, group in enumerate(groups):
            colour = ctx["colors"][index % len(ctx["colors"])]
            try:
                group_series, _, _ = _resample(
                    data[data[split] == group], date, target, {**options, "time_grain": grain}
                )
                group_series = group_series.asfreq(freq).interpolate(limit_direction="both")
                group_fit = _fit_forecast(group_series, grain, freq, horizon, model_name)
            except AnalysisError as exc:
                ctx["notes"].append(f"{group}: {exc}")
                continue
            lines["series"] += [
                {
                    "type": "line",
                    "name": str(group),
                    "data": [
                        [s, num(v)]
                        for s, v in zip(
                            c.dates(group_series.index.to_series()), group_series.values, strict=True
                        )
                    ],
                    "showSymbol": False,
                    "lineStyle": {"width": 2, "color": colour},
                    "itemStyle": {"color": colour},
                },
                {
                    "type": "line",
                    "name": str(group),
                    "data": [
                        [s, num(v)]
                        for s, v in zip(
                            c.dates(group_fit["future"].to_series()), group_fit["mean"], strict=True
                        )
                    ],
                    "showSymbol": False,
                    "lineStyle": {"width": 2.6, "type": "dashed", "color": colour},
                    "itemStyle": {"color": colour},
                },
            ]
            recent = float(group_series.iloc[-horizon:].sum())
            future_total = float(np.sum(group_fit["mean"]))
            growth_rows.append(
                (str(group), (future_total - recent) / abs(recent) if recent else float("nan"), future_total)
            )
            tables.append(
                pd.DataFrame(
                    {
                        split: group,
                        "period": group_fit["future"],
                        "forecast": group_fit["mean"],
                        "lower_80": group_fit["lo80"],
                        "upper_80": group_fit["hi80"],
                    }
                )
            )
        if growth_rows:
            charts.insert(
                1,
                _chart(
                    "forecast_by_segment",
                    f"Forecast by {split}",
                    lines,
                    ctx,
                    f"Each colour is one {split}: solid is history, dashed is its forecast.",
                ),
            )
            ranked = sorted(growth_rows, key=lambda row: -(row[1] if not math.isnan(row[1]) else -1e9))
            bars = c.base(legend=False)
            bars["tooltip"]["axisPointer"] = {"type": "shadow"}
            bars["xAxis"] = {
                "type": "value",
                "name": "Expected growth",
                "nameLocation": "middle",
                "nameGap": 28,
                "axisLabel": {"formatter": fn("percent", digits=0, scale=100)},
            }
            bars["yAxis"] = c.category_axis([row[0] for row in ranked], inverse=True)
            bars["series"] = [
                {
                    "type": "bar",
                    "data": [
                        {
                            "value": num(row[1], 4),
                            "itemStyle": {
                                "color": POSITIVE if (row[1] or 0) >= 0 else NEGATIVE,
                                "borderRadius": [0, 4, 4, 0],
                            },
                        }
                        for row in ranked
                    ],
                }
            ]
            bars["tooltip"]["valueFormatter"] = fn("percent", digits=1, scale=100)
            charts.insert(
                2,
                _chart(
                    "segment_growth",
                    f"Expected growth by {split}",
                    bars,
                    ctx,
                    f"Bars compare each {split}'s forecast with its last {horizon} periods.",
                    f"{ranked[0][0]} is expected to grow fastest ({ranked[0][1]:+.0%})."
                    if not math.isnan(ranked[0][1])
                    else "",
                ),
            )
            table = pd.concat(tables, ignore_index=True).round(4)
            metrics.insert(3, d._metric("leader", f"Fastest-growing {split}", ranked[0][0], "number"))
    mape_text = "" if math.isnan(fit["mape"]) else f", {fit['mape']:.1f}%"
    summary = [
        f"{d._model_names()[best]} gave the best backtest (average miss {fit['mae']:,.3g}{mape_text}).",
        direction or f"Forecast for the next {horizon} periods is shown with 80% and 95% ranges.",
        f"Plan for about {compact(expected)} over the next {horizon} periods (worst case "
        f"{compact(pessimistic)}, best case {compact(optimistic)}).",
    ]
    return metrics, charts, table, result, summary, comparison, best


def _change_points(values, limit: int = 3, min_size: int = 4) -> list[int]:
    """Binary segmentation on the mean: indexes where the level shifts noticeably."""
    points: list[int] = []
    segments = [(0, len(values))]
    total_sse = float(((values - values.mean()) ** 2).sum()) or 1.0
    while segments and len(points) < limit:
        best = None
        for start, end in segments:
            part = values[start:end]
            if len(part) < 2 * min_size:
                continue
            base_sse = float(((part - part.mean()) ** 2).sum())
            for split in range(min_size, len(part) - min_size):
                left, right = part[:split], part[split:]
                gain = base_sse - float(
                    ((left - left.mean()) ** 2).sum() + ((right - right.mean()) ** 2).sum()
                )
                if best is None or gain > best[0]:
                    best = (gain, start + split, (start, end))
        if best is None or best[0] / total_sse < 0.1:
            break
        _, index, segment = best
        points.append(index)
        segments.remove(segment)
        segments += [(segment[0], index), (index, segment[1])]
    return sorted(points)


def predict_trend(frame, kinds, fields, options, ctx, model_name):
    """Trend and seasonality analysis: growth, moving average, seasonal profile, change points, anomalies."""
    import numpy as np
    import pandas as pd

    date, target = fields["date"], fields["target"]
    series, grain, freq = _resample(frame, date, target, options)
    if len(series) < 6:
        raise AnalysisError(f"Trend analysis needs at least 6 periods; found {len(series)}.")
    values = series.to_numpy(dtype=float)
    window = int(options.get("window", 3))
    moving = series.rolling(window, min_periods=1).mean()
    index = np.arange(len(values))
    slope, intercept = np.polyfit(index, values, 1)
    trend_line = slope * index + intercept
    growth = series.pct_change().replace([np.inf, -np.inf], np.nan) * 100
    first, last = float(values[0]), float(values[-1])
    total_change = (last - first) / abs(first) if first else float("nan")
    average_growth = float(growth.mean()) if growth.notna().any() else float("nan")
    latest_growth = float(growth.iloc[-1]) if growth.notna().any() else float("nan")
    period = SEASON.get(grain)
    yoy = float("nan")
    if period and len(values) > period and values[-1 - period]:
        yoy = (values[-1] - values[-1 - period]) / abs(values[-1 - period])
    seasonal_strength = float("nan")
    seasonal_profile = None
    if period and len(values) >= 2 * period:
        from statsmodels.tsa.seasonal import seasonal_decompose

        decomposition = seasonal_decompose(
            series.asfreq(freq).interpolate(limit_direction="both"), model="additive", period=period
        )
        combined = (decomposition.seasonal + decomposition.resid).var()
        seasonal_strength = max(0.0, 1 - float(decomposition.resid.var() / combined)) if combined else 0.0
        residual = (decomposition.resid).to_numpy()
    else:
        residual = values - moving.to_numpy()
    residual = np.nan_to_num(residual, nan=0.0)
    mad = float(np.median(np.abs(residual - np.median(residual)))) * 1.4826 or float(np.std(residual)) or 1.0
    scores = (residual - np.median(residual)) / mad
    anomalies = np.abs(scores) > 3
    points = _change_points(values)
    if grain == "M":
        labels_ = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
        profile = series.groupby(series.index.month).mean()
        seasonal_profile = ([labels_[m - 1] for m in profile.index], profile.to_numpy())
    elif grain == "D":
        labels_ = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
        profile = series.groupby(series.index.dayofweek).mean()
        seasonal_profile = ([labels_[w] for w in profile.index], profile.to_numpy())
    elif grain == "Q":
        profile = series.groupby(series.index.quarter).mean()
        seasonal_profile = ([f"Q{q}" for q in profile.index], profile.to_numpy())
    elif grain == "W":
        profile = series.groupby(series.index.month).mean()
        seasonal_profile = (
            [
                ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"][m - 1]
                for m in profile.index
            ],
            profile.to_numpy(),
        )
    best_period, worst_period = series.idxmax(), series.idxmin()
    cv = float(np.std(values) / abs(np.mean(values))) if np.mean(values) else float("nan")
    metrics = [
        d._metric("total_change", "Change, first → last period", total_change, "percent"),
        d._metric("avg_growth", "Average growth per period", average_growth, "percent_value"),
        d._metric("latest_growth", "Latest period vs previous", latest_growth, "percent_value"),
        d._metric(
            "yoy", "Year-on-year", yoy, "percent", None, "Latest period vs the same period a year earlier."
        ),
        d._metric(
            "slope",
            "Trend per period",
            float(slope),
            "number",
            None,
            f"{target} added or lost per period on average.",
        ),
        d._metric(
            "seasonality",
            "Seasonality strength",
            seasonal_strength,
            "number",
            None,
            "0 = none, 1 = strong repeating pattern.",
        ),
        d._metric("volatility", "Volatility (coefficient of variation)", cv, "percent"),
        d._metric("anomalies", "Unusual periods", int(anomalies.sum()), "integer"),
    ]
    stamps = c.dates(series.index.to_series())
    option = c.base(zoom=len(values) > 60)
    option["xAxis"] = c.time_axis(date)
    option["yAxis"] = c.value_axis(target, scale=True)
    option["series"] = [
        {
            "type": "line",
            "name": target,
            "data": [[s, num(v)] for s, v in zip(stamps, values, strict=True)],
            "showSymbol": False,
            "lineStyle": {"width": 1.6, "opacity": 0.55},
            "itemStyle": {"color": ctx["colors"][1]},
            "markLine": {
                "silent": True,
                "symbol": "none",
                "lineStyle": {"type": "dashed", "color": "#F59E0B"},
                "label": {"formatter": "Shift", "color": "#F59E0B"},
                "data": [{"xAxis": stamps[p]} for p in points],
            },
        },
        {
            "type": "line",
            "name": f"{window}-period average",
            "data": [[s, num(v)] for s, v in zip(stamps, moving, strict=True)],
            "showSymbol": False,
            "smooth": True,
            "lineStyle": {"width": 3},
            "itemStyle": {"color": ctx["colors"][0]},
        },
        {
            "type": "line",
            "name": "Linear trend",
            "data": [[stamps[0], num(trend_line[0])], [stamps[-1], num(trend_line[-1])]],
            "showSymbol": False,
            "lineStyle": {"type": "dotted", "width": 2, "color": ctx["tokens"]["fg"]},
        },
        {
            "type": "scatter",
            "name": "Unusual period",
            "data": [[stamps[i], num(values[i])] for i in np.flatnonzero(anomalies)],
            "symbolSize": 12,
            "itemStyle": {"color": NEGATIVE},
        },
    ]
    direction = "upward" if slope > 0 else "downward"
    charts = [
        _chart(
            "trend",
            f"{target}: trend, average and unusual periods",
            option,
            ctx,
            "Thin line: actual values. Bold: moving average. Dotted: overall trend. Red dots: unusual "
            "periods. Dashed amber lines: level shifts.",
            f"The overall trend is {direction} ({compact(float(slope))} per period).",
        )
    ]
    bars = c.base(legend=False, zoom=len(values) > 60)
    bars["xAxis"] = c.time_axis(date)
    bars["yAxis"] = {
        "type": "value",
        "name": "Growth vs previous period",
        "nameLocation": "middle",
        "nameGap": 46,
        "axisLabel": {"formatter": "{value}%"},
    }
    bars["series"] = [
        {
            "type": "bar",
            "data": [
                {"value": [s, num(g, 2)], "itemStyle": {"color": POSITIVE if (g or 0) >= 0 else NEGATIVE}}
                for s, g in zip(stamps, growth, strict=True)
            ],
            "barMaxWidth": 26,
        }
    ]
    bars["tooltip"]["valueFormatter"] = fn("percent", digits=1)
    charts.append(
        _chart(
            "growth",
            "Period-over-period growth",
            bars,
            ctx,
            "Green bars grew on the previous period, red bars shrank.",
            ""
            if math.isnan(average_growth)
            else f"Average growth is {average_growth:+.1f}% per period; {int((growth > 0).sum())} of "
            f"{int(growth.notna().sum())} periods grew.",
        )
    )
    if seasonal_profile is not None:
        names, averages = seasonal_profile
        peak = names[int(np.argmax(averages))]
        profile_option = c.base(legend=False)
        profile_option["tooltip"]["axisPointer"] = {"type": "shadow"}
        profile_option["xAxis"] = c.category_axis(names)
        profile_option["yAxis"] = c.value_axis(f"Average {target}")
        profile_option["series"] = [
            {
                "type": "bar",
                "data": [
                    {
                        "value": num(v),
                        "itemStyle": {
                            "color": POSITIVE if n == peak else ctx["colors"][0],
                            "borderRadius": [4, 4, 0, 0],
                        },
                    }
                    for n, v in zip(names, averages, strict=True)
                ],
            }
        ]
        charts.append(
            _chart(
                "seasonal_profile",
                "Seasonal pattern",
                profile_option,
                ctx,
                "The average for each season position; the green bar is the strongest.",
                f"{peak} is typically the strongest period.",
            )
        )
    if period and grain in {"M", "W"} and len(values) >= period + 2:
        years = c.base()
        month_names = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
        years["xAxis"] = c.category_axis(month_names)
        years["yAxis"] = c.value_axis(target, scale=True)
        monthly = (
            series.groupby([series.index.year, series.index.month]).sum()
            if options.get("aggregation", "sum") in {"sum", "count"}
            else series.groupby([series.index.year, series.index.month]).mean()
        )
        for year in sorted({y for y, _ in monthly.index})[-5:]:
            row = monthly.loc[year]
            years["series"].append(
                {
                    "type": "line",
                    "name": str(year),
                    "data": [num(row.get(m)) if m in row.index else None for m in range(1, 13)],
                    "smooth": True,
                    "symbolSize": 5,
                }
            )
        charts.append(
            _chart(
                "year_over_year",
                "Year over year",
                years,
                ctx,
                "Each line is one year, so the same months can be compared directly.",
            )
        )
    table = pd.DataFrame(
        {
            "period": series.index,
            target: values,
            "moving_average": moving.to_numpy(),
            "growth_pct": growth.to_numpy(),
            "anomaly_score": scores,
            "unusual": anomalies,
        }
    ).round(4)
    summary = [
        f"{target} changed {total_change:+.0%} from the first to the last period; the trend is {direction}."
        if not math.isnan(total_change)
        else f"The trend is {direction}.",
        f"Best period: {str(best_period)[:10]} ({compact(float(series.max()))}); weakest: "
        f"{str(worst_period)[:10]} ({compact(float(series.min()))}).",
    ]
    if not math.isnan(yoy):
        summary.append(f"Year-on-year the latest period is {yoy:+.0%}.")
    if anomalies.any():
        summary.append(
            f"{int(anomalies.sum())} unusual period(s) stand out from the pattern — check them for "
            "one-off events."
        )
    if points:
        summary.append(
            "The level shifted around " + ", ".join(str(series.index[p])[:10] for p in points) + "."
        )
    return metrics, charts, table, table, summary, [], "robust"


def predict_drivers(frame, kinds, fields, options, ctx, model_name):
    """Key influencers: which segments push the outcome up or down, plus model-based importance."""
    import numpy as np
    import pandas as pd

    target = fields["target"]
    data = frame.dropna(subset=[target])
    data = d._sample(data, MODEL_ROWS, ctx, "analysis")
    if len(data.index) < 20:
        raise AnalysisError("At least 20 rows with an outcome are needed.")
    numeric_target = kinds.get(target) == "numeric" and d._numeric(data, target).nunique() > 10
    if numeric_target:
        outcome = d._numeric(data, target)
        positive = None
        overall = float(outcome.mean())
        measure = f"average {target}"
    else:
        labels_ = d._label(data[target])
        counts = labels_.value_counts()
        wanted = str(options.get("positive") or "")
        if wanted not in counts.index:
            preferred = [
                c_
                for c_ in counts.index
                if str(c_).lower()
                in {"1", "true", "yes", "y", "churn", "churned", "fraud", "won", "positive"}
            ]
            wanted = preferred[0] if preferred else counts.index[-1]
        positive = wanted
        outcome = (labels_ == positive).astype(float)
        overall = float(outcome.mean())
        measure = f"share of {target} = {positive}"
    features = d._inputs(data, fields, kinds, target, ctx)
    minimum = max(5, int(0.03 * len(data.index)))
    findings = []
    for column in features:
        kind = kinds.get(column, "categorical")
        if kind == "numeric":
            values = d._numeric(data, column)
            if values.nunique() < 2:
                continue
            try:
                bins = pd.qcut(values, 4, duplicates="drop")
            except ValueError:
                continue
            segments = bins.astype(str).where(bins.notna(), None)
            label_for = lambda b, col=column: f"{col} in {b}"  # noqa: E731
        elif kind == "datetime":
            continue
        else:
            segments = d._label(data[column])
            if segments.nunique() > 40:
                continue
            label_for = lambda b, col=column: f"{col} = {b}"  # noqa: E731
        grouped = (
            pd.DataFrame({"segment": segments, "outcome": outcome})
            .dropna()
            .groupby("segment")["outcome"]
            .agg(["mean", "count"])
        )
        for segment, row in grouped.iterrows():
            if row["count"] < minimum or not overall:
                continue
            lift = float(row["mean"]) / overall if overall else float("nan")
            findings.append(
                {
                    "factor": column,
                    "segment": label_for(segment),
                    "value": float(row["mean"]),
                    "rows": int(row["count"]),
                    "lift": lift,
                    "impact": abs(lift - 1) * np.sqrt(row["count"]),
                }
            )
    if not findings:
        raise AnalysisError("No segment is large enough to compare; add more rows or choose other columns.")
    ranked = sorted(findings, key=lambda item: -item["impact"])
    ups = [item for item in ranked if item["lift"] > 1][:8]
    downs = [item for item in ranked if item["lift"] < 1][:8]
    fmt = (lambda v: compact(v)) if numeric_target else (lambda v: f"{v:.0%}")  # noqa: E731
    influencers = c.base()
    influencers["tooltip"]["axisPointer"] = {"type": "shadow"}
    shown = sorted(ups + downs, key=lambda item: item["lift"])
    influencers["xAxis"] = {
        "type": "value",
        "name": f"× {measure} vs average",
        "nameLocation": "middle",
        "nameGap": 28,
    }
    influencers["yAxis"] = c.category_axis([item["segment"] for item in shown])
    influencers["series"] = [
        {
            "type": "bar",
            "name": "Lift",
            "data": [
                {
                    "value": num(item["lift"] - 1, 4),
                    "itemStyle": {"color": POSITIVE if item["lift"] > 1 else NEGATIVE, "borderRadius": 4},
                }
                for item in shown
            ],
            "label": {"show": True, "position": "right", "formatter": fn("labelValue")},
        }
    ]
    influencers["xAxis"]["axisLabel"] = {"formatter": fn("percent", digits=0, scale=100)}
    influencers["tooltip"]["valueFormatter"] = fn("percent", digits=0, scale=100)
    top = ranked[0]
    charts = [
        _chart(
            "influencers",
            f"What drives {measure}",
            influencers,
            ctx,
            "Green bars push the outcome up, red bars push it down, compared with the overall average.",
            f"When {top['segment']}, {measure} is {fmt(top['value'])} — {top['lift']:.1f}× the average "
            f"({fmt(overall)}).",
        )
    ]
    try:
        from sklearn.inspection import permutation_importance
        from sklearn.model_selection import train_test_split

        matrix = d._design_matrix(data, features, kinds, ctx)
        y = outcome if numeric_target else outcome.astype(int)
        x_train, x_test, y_train, y_test = train_test_split(matrix, y, test_size=0.25, random_state=0)
        model = d._estimator("random_forest", numeric_target).fit(x_train, y_train)
        result = permutation_importance(model, x_test, y_test, n_repeats=4, random_state=0, n_jobs=1)
        order = result.importances_mean.argsort()[::-1][:12][::-1]
        charts.append(
            _chart(
                "feature_importance",
                "Overall importance (random forest)",
                c._importance_option(
                    [str(matrix.columns[i]) for i in order],
                    result.importances_mean[order],
                    result.importances_std[order],
                    ctx,
                ),
                ctx,
                catalog.READ_GUIDE["feature_importance"],
            )
        )
    except (ValueError, KeyError) as exc:
        ctx["notes"].append(f"Importance model skipped: {str(exc)[:120]}")
    best_factor = top["factor"]
    factor_rows = [item for item in findings if item["factor"] == best_factor]
    compare = c.base(legend=False)
    compare["tooltip"]["axisPointer"] = {"type": "shadow"}
    compare["xAxis"] = c.category_axis(
        [item["segment"].split(" = ")[-1].split(" in ")[-1] for item in factor_rows], best_factor
    )
    compare["yAxis"] = c.value_axis(measure)
    compare["series"] = [
        {
            "type": "bar",
            "data": [
                {
                    "value": num(item["value"], 4),
                    "itemStyle": {
                        "color": POSITIVE if item["lift"] > 1 else ctx["colors"][0],
                        "borderRadius": [4, 4, 0, 0],
                    },
                }
                for item in factor_rows
            ],
            "markLine": {
                "silent": True,
                "symbol": "none",
                "data": [{"yAxis": num(overall, 4), "name": "Average"}],
                "lineStyle": {"type": "dashed", "color": ctx["tokens"]["muted"]},
                "label": {"formatter": "Average"},
            },
        }
    ]
    charts.append(
        _chart(
            "top_factor",
            f"{measure.capitalize()} by {best_factor}",
            compare,
            ctx,
            "Each bar is one segment; the dashed line is the overall average.",
        )
    )
    metrics = [
        d._metric("overall", f"Overall {measure}", overall, "number" if numeric_target else "percent"),
        d._metric("top_lift", "Strongest influencer", f"{top['segment']} ({top['lift']:.1f}×)", "number"),
        d._metric("factors", "Factors analysed", len(features), "integer"),
        d._metric("rows", "Rows analysed", int(len(data.index)), "integer"),
    ]
    table = pd.DataFrame(
        [
            {
                "factor": item["factor"],
                "segment": item["segment"],
                measure: round(item["value"], 4),
                "lift_vs_average": round(item["lift"], 3),
                "rows": item["rows"],
            }
            for item in ranked[:60]
        ]
    )
    summary = [f"Overall {measure} is {fmt(overall)}."]
    for item in ups[:3]:
        summary.append(
            f"When {item['segment']}, {measure} rises to {fmt(item['value'])} ({item['lift']:.1f}× average)."
        )
    for item in downs[:2]:
        summary.append(
            f"When {item['segment']}, {measure} falls to {fmt(item['value'])} ({item['lift']:.1f}× average)."
        )
    return metrics, charts, table, table, summary, [], "segments"


PREDICTORS = {
    "trend": predict_trend,
    "drivers": predict_drivers,
    "regression": predict_regression,
    "classification": predict_classification,
    "clustering": predict_clustering,
    "anomaly": predict_anomaly,
    "forecast": predict_forecast,
}


def predict(
    frame: Any, kinds: dict[str, str], request: dict[str, Any], limits: dict[str, Any]
) -> dict[str, Any]:
    task = catalog.PREDICTION_TASKS_BY_ID.get(request.get("task") or "")
    if task is None:
        raise AnalysisError("Unknown prediction task.")
    try:
        fields, options = catalog.validate_selection(
            task.slots, task.options, request.get("fields") or {}, request.get("options") or {}, kinds
        )
    except ValueError as exc:
        raise AnalysisError(str(exc)) from exc
    model_name = str(request.get("model") or task.models[0][0])
    if model_name not in dict(task.models):
        raise AnalysisError("Unknown model for this task.")
    ctx = c.context(request.get("style") or {}, limits, kinds)
    frame = d.typed(frame, kinds)
    try:
        metrics, charts, table, result, summary, comparison, model_name = PREDICTORS[task.id](
            frame, kinds, fields, options, ctx, model_name
        )
    except (KeyError, ValueError, TypeError, ZeroDivisionError, IndexError) as exc:
        if isinstance(exc, AnalysisError):
            raise
        raise AnalysisError(
            f"The model could not be trained on the chosen columns ({type(exc).__name__}: {str(exc)[:200]})."
        ) from exc
    mapping = c.label_map(frame.columns)
    if mapping:
        for chart in charts:
            c.humanize_option(chart["option"], mapping)
            chart["title"] = c.humanize_text(chart["title"], mapping, capital=True)
            explanation = chart.get("explanation") or {}
            for key in ("about", "howToRead", "insight"):
                explanation[key] = c.humanize_text(explanation.get(key, ""), mapping)
            chart["table"] = c.option_table(chart["option"])
        for metric in metrics:
            metric["label"] = c.humanize_text(metric["label"], mapping, capital=True)
            metric["help"] = c.humanize_text(metric.get("help", ""), mapping)
            if isinstance(metric.get("value"), str):
                metric["value"] = mapping.get(metric["value"], metric["value"])
        summary = [c.humanize_text(line, mapping, capital=True) for line in summary]
        ctx["notes"] = [c.humanize_text(note, mapping) for note in ctx["notes"]]
    result_bytes = result.to_csv(index=False).encode("utf-8")
    if len(result_bytes) > limits["max_result_bytes"]:
        raise AnalysisError("The prediction result exceeds the configured artifact limit.")
    if len(json.dumps(charts, separators=(",", ":"), default=str)) > limits["max_spec_bytes"] * 2:
        raise AnalysisError("The prediction charts are too large; reduce the data.")
    names = d._model_names()
    return {
        "task": task.id,
        "model": model_name,
        "modelName": names.get(model_name, model_name),
        "fields": fields,
        "options": options,
        "metrics": metrics,
        "charts": charts,
        "table": {
            "columns": [str(column) for column in table.columns],
            "rows": base.frame_preview(table, 200),
        },
        "summary": summary,
        "notes": ctx["notes"],
        "comparison": [
            {**item, "name": names.get(str(item["model"]), str(item["model"]))} for item in comparison
        ],
        "result": base64.b64encode(result_bytes).decode("ascii"),
        "resultRows": int(len(result.index)),
    }


# --- Data quality, preview and editing ----------------------------------------------------------


def quality(content: bytes, limits: dict[str, Any]) -> dict[str, Any]:
    raw = q.raw_frame(content, limits)
    return q.summarize(raw, q.scan(raw))


def preview(content: bytes, request: dict[str, Any], limits: dict[str, Any]) -> dict[str, Any]:
    """A page of rows (as typed values) plus, optionally, issue flags and an issue filter."""
    frame = d.load_frame(request, limits)
    profile, _ = d.column_kinds(frame)
    offset = max(0, int(request.get("offset") or 0))
    count = max(1, min(int(request.get("rows") or 100), 500))
    issue_filter = request.get("issue_filter") or ""
    positions = list(range(len(frame.index)))
    findings = None
    if request.get("issues") or issue_filter:
        raw = q.raw_frame(content, limits)
        findings = q.scan(raw)
        if issue_filter:
            positions = [
                int(p) for p in q.flagged_rows(findings, None if issue_filter == "issues" else issue_filter)
            ]
    page = positions[offset : offset + count]
    window = frame.iloc[page]
    payload: dict[str, Any] = {
        "profile": profile,
        "rows": base.frame_preview(window, count),
        "rowNumbers": page,
        "offset": offset,
        "total": len(positions),
        "datasetRows": int(len(frame.index)),
    }
    if findings is not None:
        payload["flags"] = {str(row): cells for row, cells in q.row_flags(findings, page).items()}
    return payload


def edit(content: bytes, request: dict[str, Any], limits: dict[str, Any]) -> dict[str, Any]:
    """Apply the user's own edits (cell values and row deletions) to the raw CSV text."""
    raw = q.raw_frame(content, limits)
    edits = request.get("edits") or []
    deletes = request.get("delete_rows") or []
    if (
        not isinstance(edits, list)
        or not isinstance(deletes, list)
        or len(edits) > 5000
        or len(deletes) > 50_000
    ):
        raise AnalysisError("Too many edits in one save; save in smaller batches.")
    rows = len(raw.index)
    for change in edits:
        if not isinstance(change, dict):
            raise AnalysisError("Each edit needs row, column and value.")
        row, column, value = change.get("row"), change.get("column"), change.get("value")
        if not isinstance(row, int) or not 0 <= row < rows:
            raise AnalysisError(f"Row {row} does not exist.")
        if column not in raw.columns:
            raise AnalysisError(f"Column {column!r} does not exist.")
        text = "" if value is None else str(value)
        if len(text) > 10_000:
            raise AnalysisError("A cell value may hold at most 10,000 characters.")
        raw.iat[row, raw.columns.get_loc(column)] = text
    drop = sorted({int(row) for row in deletes if isinstance(row, int) and 0 <= row < rows})
    if drop:
        raw = raw.drop(index=raw.index[drop])
    if raw.empty:
        raise AnalysisError("A dataset needs at least one row.")
    csv_bytes = raw.to_csv(index=False).encode("utf-8")
    frame = base.dataframe_from_bytes(csv_bytes, mime_type="text/csv", limits=limits)
    profile, _ = d.column_kinds(frame)
    return {
        "csv": base64.b64encode(csv_bytes).decode("ascii"),
        "profile": profile,
        "edited": len(edits),
        "deleted": len(drop),
    }


# --- Insights, import, organize and export -------------------------------------------------------


CLEAN_KEYS = {
    "drop_columns",
    "rename",
    "cast",
    "fill_missing",
    "drop_missing",
    "dedupe",
    "date_parts",
    "trim_text",
}


HEADER_FILL = "4F46E5"


def insights(frame: Any, profile: dict[str, Any], kinds: dict[str, str]) -> list[dict[str, Any]]:
    import numpy as np

    found: list[dict[str, Any]] = []
    rows = len(frame.index)
    if not rows:
        return found
    numeric = [c for c, k in kinds.items() if k == "numeric"]
    categorical = [c for c, k in kinds.items() if k in {"categorical", "boolean"}]
    dates = [c for c, k in kinds.items() if k == "datetime"]
    cells = rows * max(len(frame.columns), 1)
    completeness = 1 - float(frame.isna().sum().sum()) / cells
    found.append(
        {
            "id": "overview",
            "kind": "overview",
            "severity": "info",
            "title": f"{rows:,} rows × {len(frame.columns)} columns",
            "detail": f"{len(numeric)} numeric, {len(categorical)} categorical and {len(dates)} date "
            f"columns; {completeness:.1%} of cells are filled.",
        }
    )
    for entry in profile["columnProfiles"]:
        if entry["nullRatio"] >= 0.2:
            found.append(
                {
                    "id": f"missing-{entry['name']}",
                    "kind": "quality",
                    "severity": "warning",
                    "title": f"{entry['name']} is {entry['nullRatio']:.0%} empty",
                    "detail": "Fill or drop the missing values in Organize, or check the source.",
                    "chart": {"chartType": "missing_values", "fields": {}},
                }
            )
        if entry["distinct"] <= 1 and rows > 1:
            found.append(
                {
                    "id": f"constant-{entry['name']}",
                    "kind": "quality",
                    "severity": "warning",
                    "title": f"{entry['name']} has a single value",
                    "detail": "It carries no information for charts or models.",
                }
            )
    duplicates = int(frame.duplicated().sum())
    if duplicates:
        found.append(
            {
                "id": "duplicates",
                "kind": "quality",
                "severity": "warning",
                "title": f"{duplicates:,} duplicate rows",
                "detail": "Remove them in Organize if they are not genuine repeats.",
            }
        )
    if len(numeric) >= 2:
        matrix = frame[numeric[:40]].apply(lambda s: d._numeric(frame, s.name)).corr()
        pairs = matrix.where(np.triu(np.ones(matrix.shape, dtype=bool), k=1)).stack()
        for (a, b), value in pairs.reindex(pairs.abs().sort_values(ascending=False).index).head(3).items():
            if abs(value) >= 0.5:
                direction = "rise together" if value > 0 else "move in opposite directions"
                found.append(
                    {
                        "id": f"corr-{a}-{b}",
                        "kind": "correlation",
                        "severity": "positive",
                        "title": f"{a} and {b} {direction} (r = {value:.2f})",
                        "detail": "Correlation is not causation, but this pair is worth a closer look.",
                        "chart": {
                            "chartType": "scatter",
                            "fields": {"x": a, "y": b},
                            "options": {"trendline": "ols"},
                        },
                    }
                )
    for column in numeric[:30]:
        values = d._numeric(frame, column).dropna()
        if len(values) < 8:
            continue
        skew = float(values.skew())
        if abs(skew) >= 1.5:
            found.append(
                {
                    "id": f"skew-{column}",
                    "kind": "distribution",
                    "severity": "info",
                    "title": f"{column} is strongly {'right' if skew > 0 else 'left'}-skewed",
                    "detail": f"Skewness {skew:.1f}: a few {'large' if skew > 0 else 'small'} values "
                    "dominate; medians describe it better than averages.",
                    "chart": {"chartType": "histogram", "fields": {"x": column}},
                }
            )
        q1, q3 = values.quantile([0.25, 0.75])
        spread = q3 - q1
        if spread > 0:
            outliers = int(((values < q1 - 1.5 * spread) | (values > q3 + 1.5 * spread)).sum())
            if outliers and outliers / len(values) >= 0.01:
                found.append(
                    {
                        "id": f"outliers-{column}",
                        "kind": "outliers",
                        "severity": "info",
                        "title": f"{outliers:,} outliers in {column}",
                        "detail": f"{outliers / len(values):.1%} of values fall outside 1.5 × IQR.",
                        "chart": {"chartType": "box", "fields": {"y": column}},
                    }
                )
    for column in categorical[:20]:
        counts = d._label(frame[column]).value_counts()
        if 1 < len(counts) <= 50 and counts.iloc[0] / counts.sum() >= 0.5:
            found.append(
                {
                    "id": f"dominant-{column}",
                    "kind": "composition",
                    "severity": "info",
                    "title": f"“{counts.index[0]}” makes up {counts.iloc[0] / counts.sum():.0%} of {column}",
                    "detail": f"{column} has {len(counts)} values but one dominates.",
                    "chart": {"chartType": "donut", "fields": {"x": column}},
                }
            )
        if len(counts) > 0.9 * rows and rows > 20:
            found.append(
                {
                    "id": f"identifier-{column}",
                    "kind": "quality",
                    "severity": "info",
                    "title": f"{column} looks like an identifier",
                    "detail": "Almost every value is unique; it is not useful for grouping.",
                }
            )
    if dates and numeric:
        date, value = dates[0], numeric[0]
        try:
            series, _, _ = d._time_aggregate(frame, date, [value], [], "sum", "auto")
            if len(series.index) >= 4:
                y = series[value].to_numpy(dtype=float)
                slope = float(np.polyfit(np.arange(len(y)), y, 1)[0])
                change = slope * (len(y) - 1) / (abs(y.mean()) or 1)
                found.append(
                    {
                        "id": f"trend-{value}",
                        "kind": "trend",
                        "severity": "positive" if slope > 0 else "warning",
                        "title": f"{value} is trending {'up' if slope > 0 else 'down'} over {date}",
                        "detail": f"The linear trend changes it by about {change:+.0%} of its average "
                        "across the period.",
                        "chart": {"chartType": "line", "fields": {"x": date, "y": [value]}},
                    }
                )
        except AnalysisError:
            pass
    if categorical and numeric:
        group, value = next((c for c in categorical if 2 <= frame[c].nunique() <= 30), None), numeric[0]
        if group:
            means = (
                frame.assign(_v=d._numeric(frame, value))
                .groupby(d._label(frame[group]))["_v"]
                .sum()
                .sort_values(ascending=False)
            )
            if len(means) >= 2 and means.sum():
                found.append(
                    {
                        "id": f"leader-{group}-{value}",
                        "kind": "comparison",
                        "severity": "positive",
                        "title": f"{means.index[0]} leads {value} by {group}",
                        "detail": f"It accounts for {means.iloc[0] / means.sum():.0%} of the total {value}.",
                        "chart": {
                            "chartType": "bar",
                            "fields": {"x": group, "y": value},
                            "options": {"sort": "desc"},
                        },
                    }
                )
    return found[:24]


def import_table(request: dict[str, Any], limits: dict[str, Any]) -> Any:
    import pandas as pd

    content = base64.b64decode(request.get("data") or "")
    if not content:
        raise AnalysisError("The file is empty.")
    if len(content) > limits["max_bytes"]:
        raise AnalysisError("The file exceeds the configured byte limit.")
    fmt = str(request.get("format") or "csv").lower()
    try:
        if fmt == "xlsx":
            if not content.startswith(b"PK"):
                raise AnalysisError("This is not an .xlsx workbook.")
            sheet = request.get("sheet") or 0
            frame = pd.read_excel(
                io.BytesIO(content), sheet_name=sheet, engine="openpyxl", nrows=limits["max_rows"] + 1
            )
        elif fmt in {"json", "jsonl"}:
            text = _decode(content)
            if fmt == "jsonl":
                records = [json.loads(line) for line in text.splitlines() if line.strip()]
            else:
                payload = json.loads(text)
                if isinstance(payload, dict):
                    payload = next((value for value in payload.values() if isinstance(value, list)), None)
                records = payload
            if (
                not isinstance(records, list)
                or not records
                or not all(isinstance(r, dict) for r in records[:1000])
            ):
                raise AnalysisError("JSON must be an array of objects (records).")
            if len(records) > limits["max_rows"]:
                raise AnalysisError(f"Datasets may contain at most {limits['max_rows']} rows.")
            frame = pd.json_normalize(records, max_level=1)
        else:
            text = _decode(content)
            if b"\x00" in content:
                raise AnalysisError("Text files cannot contain NUL bytes.")
            separator = {"csv": ",", "tsv": "\t"}.get(fmt)
            if separator is None:
                import csv

                try:
                    separator = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|").delimiter
                except csv.Error:
                    separator = ","
            frame = pd.read_csv(
                io.StringIO(text), sep=separator, nrows=limits["max_rows"] + 1, on_bad_lines="error"
            )
    except AnalysisError:
        raise
    except (ValueError, TypeError, KeyError, OSError, pd.errors.ParserError) as exc:
        raise AnalysisError(f"The {fmt.upper()} file could not be read ({type(exc).__name__}).") from exc
    except Exception as exc:  # noqa: BLE001 - malformed workbooks raise library-specific errors
        raise AnalysisError(f"The {fmt.upper()} file could not be read ({type(exc).__name__}).") from exc
    return _normalize_columns(frame, limits)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AnalysisError("The file must use UTF-8 or Windows-1252 text encoding.")


def _normalize_columns(frame: Any, limits: dict[str, Any]) -> Any:
    frame = frame.dropna(axis=1, how="all").dropna(axis=0, how="all")
    names: list[str] = []
    for index, column in enumerate(frame.columns):
        name = str(column).strip() or f"column_{index + 1}"
        if name.startswith("Unnamed:"):
            name = f"column_{index + 1}"
        candidate, suffix = name[:120], 2
        while candidate in names:
            candidate, suffix = f"{name[:115]}_{suffix}", suffix + 1
        names.append(candidate)
    frame.columns = names
    if frame.empty or not names:
        raise AnalysisError("No data was found in the file.")
    if len(frame.columns) > limits["max_columns"]:
        raise AnalysisError(f"Datasets may contain at most {limits['max_columns']} columns.")
    if len(frame.index) > limits["max_rows"]:
        raise AnalysisError(f"Datasets may contain at most {limits['max_rows']} rows.")
    if len(frame.index) * len(frame.columns) > limits["max_cells"]:
        raise AnalysisError(f"Datasets may contain at most {limits['max_cells']} cells.")
    for column in frame.columns:
        if frame[column].dtype == object:
            frame[column] = frame[column].map(lambda v: json.dumps(v) if isinstance(v, (list, dict)) else v)
    return frame.reset_index(drop=True)


def organize(frame: Any, kinds: dict[str, str], steps: dict[str, Any]) -> Any:
    """Apply validated cleaning steps, then the shared declarative transform."""
    import pandas as pd

    if not isinstance(steps, dict):
        raise AnalysisError("Organize steps must be an object.")
    clean = steps.get("clean") or {}
    if not isinstance(clean, dict) or set(clean) - CLEAN_KEYS:
        raise AnalysisError(f"Unsupported cleaning steps: {', '.join(sorted(set(clean) - CLEAN_KEYS))}.")

    def columns_of(value: Any, label: str) -> list[str]:
        if not isinstance(value, list) or not all(isinstance(c, str) for c in value) or len(value) > 200:
            raise AnalysisError(f"{label} must list column names.")
        missing = [c for c in value if c not in frame.columns]
        if missing:
            raise AnalysisError(f"{label} names unknown columns: {', '.join(missing[:5])}.")
        return value

    frame = frame.copy()
    if clean.get("drop_columns"):
        frame = frame.drop(columns=columns_of(clean["drop_columns"], "drop_columns"))
    if clean.get("trim_text"):
        for column in frame.select_dtypes(include="object").columns:
            frame[column] = frame[column].map(lambda v: v.strip() if isinstance(v, str) else v)
    cast = clean.get("cast") or {}
    if not isinstance(cast, dict):
        raise AnalysisError("cast must map columns to number, date or text.")
    for column, kind in cast.items():
        columns_of([column], "cast")
        if kind == "number":
            frame[column] = pd.to_numeric(
                frame[column]
                .astype(str)
                .str.replace(r"[,\s$€£%]", "", regex=True)
                .replace({"": None, "nan": None}),
                errors="coerce",
            )
        elif kind == "date":
            frame[column] = (
                d._dates(frame[column]) if kinds.get(column) != "numeric" else d._dates(frame[column])
            )
        elif kind == "text":
            frame[column] = d._label(frame[column])
        else:
            raise AnalysisError("cast must be number, date or text.")
    fill = clean.get("fill_missing") or {}
    if not isinstance(fill, dict):
        raise AnalysisError("fill_missing must map columns to a strategy.")
    for column, rule in fill.items():
        columns_of([column], "fill_missing")
        strategy = rule.get("strategy") if isinstance(rule, dict) else rule
        series = frame[column]
        if strategy in {"mean", "median"}:
            numbers = pd.to_numeric(series, errors="coerce")
            frame[column] = numbers.fillna(getattr(numbers, strategy)())
        elif strategy == "mode":
            mode = series.mode(dropna=True)
            frame[column] = series.fillna(mode.iloc[0] if not mode.empty else series)
        elif strategy == "zero":
            frame[column] = series.fillna(0)
        elif strategy == "forward":
            frame[column] = series.ffill()
        elif strategy == "value":
            frame[column] = series.fillna(str((rule or {}).get("value", ""))[:200])
        else:
            raise AnalysisError("fill_missing strategy must be mean, median, mode, zero, forward or value.")
    drop_missing = clean.get("drop_missing")
    if drop_missing == "any":
        frame = frame.dropna()
    elif isinstance(drop_missing, list) and drop_missing:
        frame = frame.dropna(subset=columns_of(drop_missing, "drop_missing"))
    if clean.get("dedupe"):
        frame = frame.drop_duplicates()
    parts = clean.get("date_parts") or {}
    if not isinstance(parts, dict):
        raise AnalysisError("date_parts must map date columns to parts.")
    allowed_parts = {
        "year",
        "quarter",
        "month",
        "week",
        "weekday",
        "day",
        "hour",
        "month_name",
        "weekday_name",
    }
    for column, wanted in parts.items():
        columns_of([column], "date_parts")
        if not isinstance(wanted, list) or not set(wanted) <= allowed_parts:
            raise AnalysisError("Unsupported date part.")
        dates = d._dates(frame[column])
        for part in wanted:
            if part == "week":
                frame[f"{column}_week"] = dates.dt.isocalendar().week.astype("Int64")
            elif part == "month_name":
                frame[f"{column}_month_name"] = dates.dt.month_name()
            elif part == "weekday_name":
                frame[f"{column}_weekday_name"] = dates.dt.day_name()
            else:
                frame[f"{column}_{part}"] = getattr(dates.dt, part if part != "weekday" else "weekday")
    frame = d.base.apply_transform(frame, steps.get("transform") or {})
    rename = clean.get("rename") or {}
    if not isinstance(rename, dict) or not all(isinstance(v, str) and v.strip() for v in rename.values()):
        raise AnalysisError("rename must map columns to new, non-empty names.")
    present = {old: new.strip()[:120] for old, new in rename.items() if old in frame.columns}
    frame = frame.rename(columns=present)
    if len(set(frame.columns)) != len(frame.columns):
        raise AnalysisError("Renaming would create duplicate column names.")
    return frame.reset_index(drop=True)


def _style_sheet(sheet: Any, frame: Any) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1 and sheet.max_column:
        sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(frame.columns, start=1):
        sample = frame[column].head(200).astype(str)
        width = min(max([len(str(column))] + [len(value) for value in sample]) + 2, 48)
        sheet.column_dimensions[get_column_letter(index)].width = max(width, 8)


def _profile_table(profile: dict[str, Any]) -> Any:
    import pandas as pd

    rows = []
    for entry in profile["columnProfiles"]:
        rows.append(
            {
                "Column": entry["name"],
                "Type": entry.get("kind", ""),
                "Missing": entry["nulls"],
                "Missing %": round(entry["nullRatio"] * 100, 2),
                "Distinct": entry["distinct"],
                "Min": entry.get("min"),
                "Max": entry.get("max"),
                "Mean": entry.get("mean"),
                "Median": entry.get("median"),
                "Std. dev.": entry.get("std"),
                "Top values": ", ".join(f"{t['value']} ({t['count']})" for t in entry.get("topValues", [])),
            }
        )
    return pd.DataFrame(rows)


def _sheet_name(name: str, used: set[str]) -> str:
    cleaned = "".join("_" if ch in "[]:*?/\\" else ch for ch in name)[:28] or "Sheet"
    candidate, suffix = cleaned, 2
    while candidate.lower() in used:
        candidate, suffix = f"{cleaned[:25]}_{suffix}", suffix + 1
    used.add(candidate.lower())
    return candidate


def _excel_safe(frame: Any) -> Any:
    import pandas as pd

    safe = frame.copy()
    for column in safe.columns:
        if (
            pd.api.types.is_datetime64_any_dtype(safe[column])
            and getattr(safe[column].dt, "tz", None) is not None
        ):
            safe[column] = safe[column].dt.tz_convert(None)
        elif safe[column].dtype == object:
            safe[column] = safe[column].map(
                lambda v: f"'{v}" if isinstance(v, str) and v[:1] in "=+-@\t\r" else v
            )
    return safe


def workbook(
    sheets: list[tuple[str, Any]],
    summary: list[tuple[str, Any]],
    charts: list[dict[str, Any]] | None = None,
    graphs: bool = True,
) -> bytes:
    """Summary first, then one sheet per chart (its values and, with graphs, a native Excel chart),
    then the data sheets."""
    import pandas as pd
    from openpyxl.styles import Font
    from openpyxl.worksheet.hyperlink import Hyperlink

    output = io.BytesIO()
    used: set[str] = set()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary_frame = pd.DataFrame(summary, columns=["Item", "Value"])
        summary_name = _sheet_name("Summary", used)
        summary_frame.to_excel(writer, sheet_name=summary_name, index=False)
        summary_sheet = writer.sheets[summary_name]
        _style_sheet(summary_sheet, summary_frame)
        summary_sheet["A1"].font = Font(bold=True, color="FFFFFF")
        summary_sheet.column_dimensions["B"].width = 80
        for title, frame in sheets:
            name = _sheet_name(title, used)
            safe = _excel_safe(frame)
            safe.to_excel(writer, sheet_name=name, index=False)
            _style_sheet(writer.sheets[name], safe)
        book = writer.book
        for index, chart in enumerate(charts or []):
            label = f"Chart {index + 1}"
            if chart.get("error"):
                summary_sheet.append(
                    [label, f"{chart.get('title') or 'Chart'} — not included: {chart['error']}"]
                )
                continue
            name = _sheet_name(f"{label} {chart.get('title') or ''}".strip(), used)
            table = chart.get("table") or c.option_table(chart.get("option") or {})
            drew = x.add_chart_sheet(book, name, chart, table, graphs)
            book.move_sheet(name, offset=-(len(book.sheetnames) - 2 - index))
            text = str(chart.get("title") or label)
            if chart.get("subtitle"):
                text += f" ({chart['subtitle']})"
            if graphs and not drew:
                text += " — values only (no Excel equivalent; use the HTML report for the interactive chart)"
            summary_sheet.append([label, text])
            cell = summary_sheet.cell(row=summary_sheet.max_row, column=2)
            cell.hyperlink = Hyperlink(
                ref=cell.coordinate, location="'{}'!A1".format(name.replace("'", "''"))
            )
            cell.font = Font(color="4F46E5", underline="single")
            for stat in chart.get("stats") or []:
                summary_sheet.append([f"  {stat.get('label')}", str(stat.get("value"))])
    return output.getvalue()


def chart_entries(
    items: list[dict[str, Any]], frames: list[tuple[Any, dict[str, str]]], limits: dict[str, Any]
) -> list[dict[str, Any]]:
    """Charts for a workbook: ready ECharts options, or specs rendered here from a dataset."""
    entries: list[dict[str, Any]] = []
    for item in items[:40]:
        entry = {
            "title": str(item.get("title") or "")[:200],
            "subtitle": str(item.get("subtitle") or "")[:300],
            "note": str(item.get("note") or "")[:1000],
            "stats": item.get("stats") or [],
        }
        if isinstance(item.get("option"), dict):
            entries.append({**entry, "option": item["option"]})
            continue
        index = item.get("dataset", 0)
        if not isinstance(index, int) or not 0 <= index < len(frames):
            continue
        frame, kinds = frames[index]
        try:
            rendered = c.render(
                frame,
                kinds,
                {
                    "chart_type": item.get("chart_type"),
                    "fields": item.get("fields") or {},
                    "options": item.get("options") or {},
                    "style": {**(item.get("style") or {}), "title": entry["title"]},
                },
                limits,
            )
        except AnalysisError as exc:
            entries.append({**entry, "error": str(exc)[:200]})
            continue
        explanation = rendered.get("explanation") or {}
        entries.append(
            {
                **entry,
                "title": rendered["title"],
                "note": entry["note"] or explanation.get("insight", ""),
                "option": rendered["option"],
                "table": rendered["table"],
                "stats": entry["stats"] or rendered.get("stats") or [],
            }
        )
    return entries


def export_dataset(
    frame: Any,
    kinds: dict[str, str],
    profile: dict[str, Any],
    request: dict[str, Any],
    limits: dict[str, Any],
) -> dict[str, Any]:
    fmt = request.get("format", "csv")
    name = str(request.get("name") or "Dataset")[:120]
    if fmt == "csv":
        content = frame.to_csv(index=False).encode("utf-8")
        mime = "text/csv"
    elif fmt == "json":
        content = frame.to_json(orient="records", date_format="iso", force_ascii=False).encode("utf-8")
        mime = "application/json"
    elif fmt == "xlsx":
        typed_frame = d.typed(frame, kinds)
        found = insights(typed_frame, profile, kinds)
        summary = [("Dataset", name), ("Rows", len(frame.index)), ("Columns", len(frame.columns))]
        summary += [
            (f"Insight {i + 1}", f"{item['title']}. {item.get('detail', '')}") for i, item in enumerate(found)
        ]
        graphs = bool(request.get("graphs"))
        charts = chart_entries(request.get("charts") or [], [(frame, kinds)], limits) if graphs else []
        content = workbook(
            [("Data", typed_frame), ("Profile", _profile_table(profile))], summary, charts, graphs
        )
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        raise AnalysisError("Unsupported export format.")
    if len(content) > limits["max_result_bytes"] * 2:
        raise AnalysisError("The export exceeds the configured size limit.")
    return {"content": base64.b64encode(content).decode("ascii"), "mimeType": mime}


def report_workbook(request: dict[str, Any], limits: dict[str, Any]) -> dict[str, Any]:
    """A workbook from reports, Studio pages and predictions: summary, chart sheets, tables, data."""
    sheets: list[tuple[str, Any]] = []
    for item in (request.get("tables") or [])[:4]:
        frame = d.base.dataframe_from_bytes(
            base64.b64decode(item["data"]), mime_type="text/csv", limits=limits
        )
        sheets.append((str(item.get("name") or "Table")[:28], frame))
    frames: list[tuple[Any, dict[str, str]]] = []
    for item in (request.get("datasets") or [])[:8]:
        frame = d.base.dataframe_from_bytes(
            base64.b64decode(item["data"]), mime_type="text/csv", limits=limits
        )
        profile, kinds = d.column_kinds(frame)
        frames.append((frame, kinds))
        if item.get("sheets", True):
            label = str(item.get("name") or "Dataset")[:20]
            sheets.append((f"{label} data", d.typed(frame, kinds)))
            sheets.append((f"{label} profile", _profile_table(profile)))
    summary = [(str(k)[:200], str(v)[:2000]) for k, v in (request.get("summary") or [])][:500]
    graphs = bool(request.get("graphs", True))
    charts = chart_entries(request.get("charts") or [], frames, limits)
    content = workbook(sheets, summary, charts, graphs)
    return {
        "content": base64.b64encode(content).decode("ascii"),
        "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }


def _finite(value: Any) -> bool:
    try:
        json.dumps(value, allow_nan=False, default=str)
        return True
    except ValueError:
        return False


def _clean_nan(value: Any) -> Any:
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, dict):
        return {k: _clean_nan(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_clean_nan(v) for v in value]
    return value


# --- Dashboards ----------------------------------------------------------------------------------


def query_table(frame: Any, kinds: dict[str, str]) -> dict[str, Any]:
    """The dataset as clean, typed CSV for the Postgres query store (columns renamed c0…cN)."""
    import pandas as pd

    data = d.typed(frame, kinds)
    out = pd.DataFrame(index=data.index)
    columns: list[dict[str, Any]] = []
    for index, name in enumerate(data.columns):
        kind = kinds.get(name, "categorical")
        column = f"c{index}"
        series = data[name]
        if kind == "numeric":
            out[column] = pd.to_numeric(series, errors="coerce")
        elif kind == "datetime":
            out[column] = pd.to_datetime(series, errors="coerce").dt.strftime("%Y-%m-%d %H:%M:%S")
        elif kind == "boolean":
            text = series.astype("string").str.strip().str.lower()
            out[column] = text.map(
                lambda v: (
                    "true"
                    if v in {"true", "1", "yes", "y"}
                    else "false"
                    if v in {"false", "0", "no", "n"}
                    else None
                )
            )
        else:
            out[column] = series.map(
                lambda v: None if v is None or (isinstance(v, float) and math.isnan(v)) else str(v)
            )
        columns.append({"name": str(name), "column": column, "kind": kind})
    out["_row"] = range(len(out.index))
    content = out.to_csv(index=False, na_rep="").encode("utf-8")
    return {
        "csv": base64.b64encode(content).decode("ascii"),
        "columns": columns,
        "rows": int(len(out.index)),
    }


def dashboard_visuals(request: dict[str, Any], limits: dict[str, Any]) -> dict[str, Any]:
    """Build a batch of dashboard visuals. Results come pre-aggregated from Postgres, or are
    aggregated here with pandas when the request carries the dataset itself."""
    frame = None
    kinds = dict(request.get("kinds") or {})
    if request.get("data"):
        frame = d.load_frame(request, limits)
        kinds = d.column_kinds(frame)[1]
        frame = d.typed(frame, kinds)
    ctx = vv.Ctx(request["theme"], request.get("categoryOrder") or {})
    tokens = {"grid": ctx.s["gridline"], "bg": ctx.s["tile"]}
    geo = {
        "world_names": c.world_names,
        "continents": c.world_continents,
        "geo_base": lambda scope: c._geo_base({"tokens": tokens, "colors": ctx.colors}, scope),
        "render": lambda table, table_kinds, spec: c.render(table, table_kinds, spec, limits),
        "kinds": kinds,
    }
    visuals: list[dict[str, Any]] = []
    computed: dict[str, Any] = {}
    for item in request.get("items") or []:
        visual = item.get("visual") or {}
        try:
            results = dict(item.get("results") or {})
            for role, query in (item.get("queries") or {}).items():
                if role not in results:
                    if frame is None:
                        raise AnalysisError("The data for this visual is not available.")
                    results[role] = vq.aggregate(frame, query, kinds)
                    if item.get("keys", {}).get(role):
                        computed[item["keys"][role]] = results[role]
            visuals.append(vv.build(visual, item.get("queries") or {}, results, ctx, **geo))
        except AnalysisError as exc:
            visuals.append(
                {
                    "id": visual.get("id"),
                    "type": visual.get("type"),
                    "status": "error",
                    "message": str(exc)[:300],
                }
            )
        except (KeyError, ValueError, TypeError, IndexError, ZeroDivisionError) as exc:
            visuals.append(
                {
                    "id": visual.get("id"),
                    "type": visual.get("type"),
                    "status": "error",
                    "message": f"This visual could not be drawn ({type(exc).__name__}).",
                }
            )
    return {"visuals": visuals, "results": computed}


# --- Process entry point -----------------------------------------------------------------------


def handle(request: dict[str, Any]) -> dict[str, Any]:
    operation = request["op"]
    limits = request["limits"]
    if operation == "viz_import":
        frame = import_table(request, limits)
        profile, _ = d.column_kinds(frame)
        content = frame.to_csv(index=False).encode("utf-8")
        return {
            "csv": base64.b64encode(content).decode("ascii"),
            "profile": profile,
            "preview": base.frame_preview(frame, 20),
        }
    if operation == "viz_visuals":
        return dashboard_visuals(request, limits)
    if operation in {"viz_report_xlsx", "viz_workbook"}:
        return report_workbook(request, limits)
    content = base64.b64decode(request.get("data") or "")
    if operation == "viz_quality":
        return quality(content, limits)
    if operation == "viz_edit":
        return edit(content, request, limits)
    if operation == "viz_preview":
        return preview(content, request, limits)
    frame = d.load_frame(request, limits)
    profile, kinds = d.column_kinds(frame)
    if operation == "viz_table":
        return query_table(frame, kinds)
    if operation == "viz_render":
        return c.render(frame, kinds, request, limits)
    if operation == "viz_predict":
        return predict(frame, kinds, request, limits)
    if operation == "viz_insights":
        return {"profile": profile, "insights": insights(d.typed(frame, kinds), profile, kinds)}
    if operation == "viz_organize":
        result = organize(frame, kinds, request.get("steps") or {})
        csv_bytes = result.to_csv(index=False).encode("utf-8")
        if len(csv_bytes) > limits["max_result_bytes"]:
            raise AnalysisError("The organized dataset exceeds the configured size limit.")
        result_profile, _ = d.column_kinds(result)
        return {
            "profile": result_profile,
            "preview": base.frame_preview(result, int(limits["preview_rows"])),
            "csv": base64.b64encode(csv_bytes).decode("ascii"),
        }
    if operation == "viz_export":
        return export_dataset(frame, kinds, profile, request, limits)
    raise AnalysisError(f"Unsupported operation {operation!r}.")


def main() -> int:
    import os

    limits = json.loads(os.environ.get("JT_ANALYTICS_LIMITS", "{}"))
    base._apply_resource_limits(limits)
    # From here on the process cannot reach the network, start programs or touch files outside
    # its private temporary directory (see viz_guard.py).
    _load("jt_viz_guard", "viz_guard.py").install(engine_dirs=(str(_HERE.parent / "analytics"),))
    try:
        request = json.loads(sys.stdin.buffer.read())
        response: dict[str, Any] = {"ok": True, **handle(request)}
    except AnalysisError as exc:
        response = {"ok": False, "kind": "analysis", "error": str(exc)[:2000]}
    except MemoryError:
        response = {"ok": False, "kind": "resources", "error": "The operation exceeded its memory limit."}
    except Exception as exc:  # noqa: BLE001 - reported to the parent without a traceback
        response = {"ok": False, "kind": "internal", "error": f"{type(exc).__name__}: {str(exc)[:300]}"}
    sys.stdout.write(
        json.dumps(response, default=str, allow_nan=False)
        if _finite(response)
        else json.dumps(_clean_nan(response), default=str)
    )
    sys.stdout.flush()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
